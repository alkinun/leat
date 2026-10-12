"""The workspace, its tools and the sandbox. The sandbox's tests need bubblewrap, and those of the
documents it reads LEAT_SANDBOX, an environment made of leat/agent/sandbox.txt."""

import datetime
import http.client
import json
import os
import shutil
import urllib.error
import urllib.request
from collections.abc import Iterator
from pathlib import Path

import pytest

from leat.agent import agent as agent_module
from leat.agent import workspace as workspace_module
from leat.agent.agent import Agent, NotFound
from leat.agent.client import Client
from leat.agent.context import message as _api
from leat.agent.store import Store
from leat.agent.tools import files
from leat.agent.workspace import Workspace
from tests.test_agent import call as calling
from tests.test_agent import cookie, ended, serving, until


def _sandboxes() -> bool:
    # whether bubblewrap runs here: some systems forbid its user namespaces
    if shutil.which("bwrap") is None or shutil.which("prlimit") is None:
        return False
    return Workspace(Path("/tmp")).run("print(1)", timeout=30).output.strip() == "1"


SANDBOX = os.environ.get("LEAT_SANDBOX")
sandboxed = pytest.mark.skipif(not _sandboxes(), reason="needs bubblewrap, and its namespaces")
documents = pytest.mark.skipif(not SANDBOX, reason="needs LEAT_SANDBOX, the sandbox's environment")


@pytest.fixture
def workspace(tmp_path) -> Workspace:
    return Workspace(tmp_path / "workspace", Path(SANDBOX) if SANDBOX else None)


def test_paths(workspace, tmp_path):
    # a name is the workspace's, an absolute one too; one out of it is refused, by a link too
    assert workspace.path("notes/a.txt") == workspace.root / "notes" / "a.txt"
    assert workspace.path("/etc/passwd") == workspace.root / "etc" / "passwd"
    with pytest.raises(ValueError, match="outside the workspace"):
        workspace.path("../leat.db")
    (workspace.root / "out").symlink_to(tmp_path)
    with pytest.raises(ValueError, match="outside the workspace"):
        workspace.path("out/leat.db")


def test_files(workspace, tmp_path):
    (workspace.root / "notes").mkdir()
    (workspace.root / "notes" / "a.txt").write_text("a")
    (workspace.root / ".hidden").write_text("b")
    (workspace.root / "b.txt").write_text("bb")
    os.utime(workspace.root / "notes" / "a.txt", (0, 0))
    (tmp_path / "key").write_text("secret")
    (workspace.root / "key").symlink_to(tmp_path / "key")  # out of it
    assert [(f["name"], f["size"]) for f in workspace.files()] == [("b.txt", 2), ("notes/a.txt", 1)]
    # a name for an upload: at the top, free, and of no path
    assert workspace.free("b.txt") == "b (2).txt" and workspace.free("../../x.pdf") == "x.pdf"
    assert workspace.free(".bashrc") == "bashrc" and workspace.free("") == "file"
    # a name past what a file system takes: refused, or for an upload, its start and its kind
    with pytest.raises(ValueError, match="255 bytes at most"):
        workspace.path("é" * 128)
    long = workspace.free("é" * 200 + ".txt")
    assert long.endswith("é.txt") and len(long.encode()) <= 247
    (workspace.root / long).write_text("x")
    numbered = workspace.free("é" * 200 + ".txt")
    assert numbered == long.removesuffix(".txt") + " (2).txt" and len(numbered.encode()) <= 255
    workspace.delete(long)
    workspace.delete("b.txt")
    with pytest.raises(FileNotFoundError):
        workspace.delete("b.txt")
    # a file deleted between being found and looked at, as by code a call runs meanwhile, is gone
    stat, gone, looks = Path.stat, workspace.root / "notes" / "a.txt", []

    def deleting(path: Path, **kwargs) -> os.stat_result:  # at the second look, after is_file's
        if path == gone:
            looks.append(path)
            if len(looks) == 2:
                gone.unlink()
        return stat(path, **kwargs)

    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(Path, "stat", deleting)
        assert workspace.files() == []


def test_read_write_edit(workspace):
    files.write(workspace, "notes/plan.md", "# Plan\n\nBuy bread.\n")
    assert files.read(workspace, "notes/plan.md").content == "# Plan\n\nBuy bread.\n"
    assert files.read(workspace, ".").content == "notes/"
    today = datetime.date.today()
    assert files.read(workspace, "notes").content == f"plan.md (19 B, changed {today})"
    result = files.edit(workspace, "notes/plan.md", "bread", "milk")
    assert result.info == {"files": ["notes/plan.md"]}
    assert (workspace.root / "notes" / "plan.md").read_text() == "# Plan\n\nBuy milk.\n"
    with pytest.raises(ValueError, match="is not in"):
        files.edit(workspace, "notes/plan.md", "bread", "milk")
    files.write(workspace, "twice.txt", "a a")
    with pytest.raises(ValueError, match="is in 2 places of"):
        files.edit(workspace, "twice.txt", "a", "b")
    # a long file in parts, saying where the next begins
    files.write(workspace, "long.txt", "x" * (files.READ + 10))
    first = files.read(workspace, "long.txt").content
    assert first.endswith(f"read on from start={files.READ})")
    assert files.read(workspace, "long.txt", files.READ).content == "x" * 10
    # what is not text
    (workspace.root / "photo.jpg").write_bytes(b"\xff\xd8")
    (workspace.root / "photo.heic").write_bytes(b"\0")
    (workspace.root / "blob").write_bytes(b"\0\1\2")
    seen = files.read(workspace, "photo.jpg")  # for the agent to show the model
    assert (seen.content, seen.info) == (
        "The image photo.jpg.",
        {"file": "photo.jpg", "images": ["photo.jpg"]},
    )
    with pytest.raises(ValueError, match="convert it to a PNG"):
        files.read(workspace, "photo.heic")
    with pytest.raises(ValueError, match="not a file of text"):
        files.read(workspace, "blob")
    with pytest.raises(FileNotFoundError):
        files.read(workspace, "nothing.txt")


def test_skills(workspace):
    # each listed, and read from skills/, but nothing beside them
    found = dict(files.skills())
    assert set(found) == {f"skills/{k}/SKILL.md" for k in ("docx", "xlsx", "pptx", "pdf")}
    assert "python-docx" in files.read(workspace, "skills/docx/SKILL.md").content
    with pytest.raises(ValueError, match="outside the skills"):
        files.read(workspace, "skills/../tools/files.py")


@sandboxed
def test_run(workspace):
    # what the code prints, and the files it makes; the workspace alone, without the network
    code = """
import os, socket
open("made.txt", "w").write("hi")
print(sorted(os.listdir("/")), os.path.exists(os.path.expanduser("~/.ssh")))
try:
    socket.create_connection(("1.1.1.1", 80), timeout=2)
except OSError:
    print("no network")
"""
    (workspace.root / "json.py").write_text("raise SystemExit('a library taken')")
    result = files.run(workspace, code + "import json\n")
    assert "no network" in result.content and "False" in result.content
    assert result.info == {"files": ["made.txt"], "status": 0}
    assert result.content.endswith("Files made or changed: made.txt")
    assert "It failed, with exit status 1." in files.run(workspace, "1 / 0").content
    assert workspace.run("while True: pass", timeout=1).status is None
    # a loop of prints, read as it comes, of which the start and the end are kept
    printed = workspace.run("while True: print('x' * 1000)", timeout=1)
    assert printed.status is None and "bytes left out" in printed.output
    assert len(printed.output) < workspace_module.OUTPUT + 100


@documents
def test_documents(workspace):
    # what the libraries make, read back
    files.run(
        workspace,
        """
import docx, openpyxl
d = docx.Document(); d.add_heading("Plan", 0); d.add_paragraph("Buy milk."); d.save("plan.docx")
wb = openpyxl.Workbook(); wb.active.append(["Rent", 900, "=B1*12"]); wb.save("budget.xlsx")
""",
    )
    assert files.read(workspace, "plan.docx").content == "# Plan\n\nBuy milk.\n\n"
    # a Word document's header and footer, as its letterhead, around its body
    files.run(workspace, """
import docx
d = docx.Document()
d.sections[0].header.paragraphs[0].text = "Becker & Partner"
d.sections[0].footer.paragraphs[0].text = "80802 München"
d.save("Briefkopf.docx")
""")  # fmt: skip
    letterhead = "Header: Becker & Partner\n\nFooter: 80802 München\n\n"
    assert files.read(workspace, "Briefkopf.docx").content == letterhead
    assert files.read(workspace, "budget.xlsx").content == (
        "## Sheet\n\n\n| Rent | 900 | =B1*12 |\n| --- | --- | --- |\n\n")  # fmt: skip
    # a long one, in parts, whole
    files.run(workspace, """
import docx
d = docx.Document()
for i in range(400):
    d.add_paragraph(f"Line {i}: " + "words " * 10)
d.save("long.docx")
""")  # fmt: skip
    first = files.read(workspace, "long.docx").content
    assert first.endswith(f"read on from start={files.READ})")
    whole, read = "", first
    while "read on from start=" in read:
        whole += read.rsplit("\n\n(characters", 1)[0]
        read = files.read(workspace, "long.docx", int(read.rsplit("=", 1)[1][:-1])).content
    lines = (whole + read).split("\n\n")
    assert lines[:-1] == [f"Line {i}: {'words ' * 10}" for i in range(400)]
    with pytest.raises(ValueError, match="could not be read"):
        files.write(workspace, "broken.pdf", "not a PDF")
        files.read(workspace, "broken.pdf")


@pytest.fixture
def agent(workspace, tmp_path) -> Agent:
    return Agent(Store(tmp_path / "leat.db"), Client("http://127.0.0.1:9"), [], workspace)


def test_attached(agent, workspace):
    # the files attached to a message, named for the model, which must be in its space
    files.write(agent.space(), "plan.md", "Buy milk.")
    with pytest.raises(NotFound):
        agent.send(None, "Read it", attached=["nothing.md"])
    message = {"role": "user", "content": "Read it", "info": {"files": ["plan.md"]}}
    assert _api(message)["content"] == "Read it\n\n(Attached, in the workspace: plan.md)"
    with agent.events.watch() as events:
        agent.files_changed()
        event = events.get(timeout=1)
        assert event["files"][0]["name"] == "plan.md" and event["to"] is None
        agent.files_changed()  # unchanged: nothing
        assert events.empty()


@pytest.mark.parametrize("vision", [True, False])
def test_images_seen(engine, workspace, tmp_path, vision):
    # an image the user attaches, and one a tool reads, the model sees as data: URLs before their
    # text, if it sees images; if not, the user's are named, and a tool's said to be unseen
    engine.vision = vision
    store = Store(tmp_path / "leat.db")
    agent = Agent(store, Client(engine.url), files.tools(), workspace)
    (agent.space().root / "cat.png").write_bytes(b"\x89PNG cat")
    engine.replies.put([{"tool_calls": [calling("read", {"path": "cat.png"})]}])
    engine.replies.put([{"content": "A cat."}])
    with agent.events.watch() as events:
        agent.send(None, "What is it?", attached=["cat.png"])
        until(events, ended)
    first, second = engine.requests
    user, tool = first["messages"][1]["content"], second["messages"][-1]["content"]
    url = {"type": "image_url", "image_url": {"url": "data:image/png;base64,iVBORyBjYXQ="}}
    if vision:
        assert user[0] == url and user[1]["text"].endswith("(Attached, in the workspace: cat.png)")
        assert tool == [url, {"type": "text", "text": "The image cat.png."}]
    else:
        assert user.endswith("(Attached, in the workspace: cat.png)")
        assert tool == "The image cat.png. You cannot see it: the model takes no images."
    assert "images" not in second["messages"][1]


def test_system_prompt(agent, tmp_path):
    # the workspace, when the agent has one, with its tools' rules: run's, of the skills
    from leat.agent.agent import _system

    with_workspace, without = (_system(w, tools=files.tools())["content"] for w in (True, False))
    assert "with run. The skills:\n  - skills/docx/SKILL.md: Word documents" in with_workspace
    assert "workspace" not in without and "skills" not in without
    assert "skills" not in _system(True)["content"]  # without run


@pytest.fixture
def server(agent) -> Iterator[str]:
    with serving(agent) as url:
        yield url


def call(url: str, method: str = "GET", data: bytes | None = None):
    try:
        with urllib.request.urlopen(urllib.request.Request(url, data, method=method)) as r:
            return r.status, r.headers, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.headers, e.read()


def test_api_files(server, agent, monkeypatch):
    # uploads to the person's own files, by a free name, downloads, which a page cannot act in,
    # and deletes
    status, _, body = call(f"{server}/api/files/My%20notes.txt", "PUT", b"Buy milk.")
    assert status == 200 and json.loads(body) == {"name": "My notes.txt"}
    assert json.loads(call(f"{server}/api/files/My%20notes.txt", "PUT", b"x")[2])["name"] == (
        "My notes (2).txt")  # fmt: skip
    status, headers, body = call(f"{server}/files/My%20notes.txt")
    assert status == 200 and body == b"Buy milk."
    assert headers["Content-Security-Policy"] == "sandbox"
    assert headers["Content-Disposition"] == "inline; filename*=UTF-8''My%20notes.txt"
    assert agent.space(None, 1).path("My notes.txt").read_bytes() == b"Buy milk."
    files.write(agent.space(None, 1), "page.html", "<script>fetch('/api/events')</script>")
    assert call(f"{server}/files/page.html")[1]["Content-Disposition"].startswith("attachment")
    for path in ("/files/..%2Fleat.db", "/files/nothing.txt"):
        assert call(f"{server}{path}")[0] == 404
    assert call(f"{server}/api/files/My%20notes.txt", "DELETE")[0] == 200
    assert call(f"{server}/api/files/My%20notes.txt", "DELETE")[0] == 404
    # the file deleted kept as an earlier version, and restored
    versions = json.loads(call(f"{server}/api/history")[2])["versions"]
    assert [(v["name"], v["deleted"]) for v in versions] == [("My notes.txt", True)]
    restored = json.dumps({"name": "My notes.txt", "version": versions[0]["version"]}).encode()
    assert call(f"{server}/api/history", "POST", restored)[0] == 200
    assert agent.space(None, 1).path("My notes.txt").read_bytes() == b"Buy milk."
    unknown = json.dumps({"name": "My notes.txt", "version": "1"}).encode()
    assert call(f"{server}/api/history", "POST", unknown)[0] == 404
    assert agent.store.activity()[0]["action"] == "restored"
    for length in ("-1", "many"):  # a length there cannot be
        bad = urllib.request.Request(f"{server}/api/files/x.txt", b"x", method="PUT")
        bad.add_unredirected_header("Content-Length", length)
        with pytest.raises(urllib.error.HTTPError, match="400"):
            urllib.request.urlopen(bad)
    # a body in chunks, of no length, which would make an empty file
    chunked = urllib.request.Request(f"{server}/api/files/c.txt", iter([b"x"]), method="PUT")
    chunked.add_unredirected_header("Transfer-Encoding", "chunked")
    with pytest.raises(urllib.error.HTTPError, match="400"):
        urllib.request.urlopen(chunked)
    assert not agent.space(None, 1).path("c.txt").exists()
    monkeypatch.setattr("leat.agent.server.UPLOAD", 4)
    assert call(f"{server}/api/files/big.bin", "PUT", b"12345")[0] == 413
    assert call(f"{server}/files")[0] == 200  # the app's page of them


def test_workflow_on_arrival(server, agent, engine, monkeypatch):
    # a workflow run as files arrive in its project, uploaded or synced, once they settle, the
    # files attached, those uploaded together in one run; not by files Leat makes, nor again by
    # those it ran for; one gone since, not attached
    monkeypatch.setattr(agent_module, "SETTLE", 0)
    agent.engine = Client(engine.url)
    project = agent.add_project("Payroll", 1)["id"]
    flow = agent.add_workflow("New scans", "Read each new scan.", 1, project)
    agent.arrived(project, 1, ["early.txt"])  # before any workflow runs on arrival: not kept
    agent.schedule_workflow(flow["id"], "files", 1)
    for name in ("a.txt", "b.txt"):
        assert call(f"{server}/api/projects/{project}/files/{name}", "PUT", b"x")[0] == 200
    files.write(agent.space(project, 1), "made.txt", "by a tool")  # made by Leat: not arrived
    engine.status = "unloaded"
    agent.run_due()  # no model to run it: the files kept for when there is
    assert agent.conversations(1) == []
    engine.status = "loaded"
    engine.replies.put([{"content": "Read."}])
    with agent.events.watch() as events:
        agent.run_due()
        until(events, ended)
    (run,) = agent.conversations(1)
    assert agent.store.messages(run["id"])[1]["info"]["files"] == ["a.txt", "b.txt"]
    agent.run_due()
    assert len(agent.conversations(1)) == 1  # nothing new arrived
    agent.arrived(project, None, ["gone.txt"])  # synced, then deleted before it settled
    agent.run_due()
    assert len(agent.conversations(1)) == 1


def test_spaces(agent, workspace, engine, tmp_path):
    # each person's own files, and each project's, apart: a conversation's tools work in its
    # space alone, and a person reaches the spaces they see alone
    owner, ada = agent.store.add_person("Alkın")["id"], agent.store.add_person("Ada")["id"]
    project = agent.add_project("Yılmaz Ltd", owner)["id"]
    files.write(agent.space(None, owner), "mine.txt", "mine")
    files.write(agent.space(project, owner), "contract.txt", "the contract")
    assert agent.space(None, owner).root == workspace.root / "people" / str(owner)
    assert agent.space(project, owner).root == workspace.root / "projects" / project
    with pytest.raises(NotFound):
        agent.space(project, ada)
    with pytest.raises(NotFound):
        agent.send(None, "Read it", attached=["mine.txt"], person=owner, project=project)
    store = agent.store
    agent = Agent(store, Client(engine.url), files.tools(), workspace)
    engine.replies.put([{"tool_calls": [calling("read", {"path": "contract.txt"})]}])
    engine.replies.put([{"content": "Read."}])
    with agent.events.watch() as events:
        id = agent.send(None, "Read it", attached=["contract.txt"], person=owner, project=project)
        until(events, ended)
    assert store.messages(id)[-2]["content"] == "the contract"
    # deleted, a project takes its files
    agent.delete_project(project, owner)
    assert not (workspace.root / "projects" / project).exists()


def test_settled(tmp_path):
    # the files of the workspace before it had spaces, and those of no one's, become the owner's
    workspace = Workspace(tmp_path / "workspace")
    (workspace.root / "old.txt").write_text("old")
    (workspace.root / ".web").mkdir()
    store = Store(tmp_path / "leat.db")
    agent = Agent(store, Client("http://127.0.0.1:9"), [], workspace)
    assert [f["name"] for f in agent.space().files()] == ["old.txt"]
    assert (workspace.root / "people" / "0" / ".web").is_dir()
    owner = store.add_person("Alkın")["id"]
    (workspace.root / "people" / str(owner)).mkdir(parents=True)
    (workspace.root / "people" / str(owner) / "old.txt").write_text("kept")
    agent.settle()
    assert sorted(f["name"] for f in agent.space(None, owner).files()) == ["old (2).txt", "old.txt"]
    assert not (workspace.root / "people" / "0").exists()


def test_api_project_files(server, agent):
    # a project's files, uploaded, downloaded and deleted by those who see it alone
    project = agent.add_project("Payroll", 1)["id"]
    base = f"{server}/api/projects/{project}/files"
    assert json.loads(call(f"{base}/March.xlsx", "PUT", b"cells")[2])["name"] == "March.xlsx"
    assert agent.space(project, 1).path("March.xlsx").read_bytes() == b"cells"
    assert call(f"{server}/projects/{project}/files/March.xlsx")[2] == b"cells"
    assert call(f"{server}/files/March.xlsx")[0] == 404  # not the person's own
    assert call(f"{server}/projects/{project}")[0] == 200  # the app's page of it
    ada = agent.store.add_person("Ada")["id"]
    other = agent.add_project("Hers", ada)["id"]
    for method, url in [("PUT", f"{server}/api/projects/{other}/files/x.txt"),
                        ("GET", f"{server}/projects/{other}/files/x.txt"),
                        ("DELETE", f"{server}/api/projects/{other}/files/x.txt")]:  # fmt: skip
        assert call(url, method, b"x" if method == "PUT" else None)[0] == 404
    assert call(f"{base}/March.xlsx", "DELETE")[0] == 200
    assert agent.space(project, 1).files() == []


def test_api_project_files_of_members(server, agent):
    # a project's files, of one shared with chosen people, seen by those, and by its maker alone,
    # each told of them: uploaded, downloaded, deleted and their earlier versions
    ada, ben = agent.store.add_person("Ada")["id"], agent.store.add_person("Ben")["id"]
    project = agent.add_project("Hers", ada)["id"]
    agent.files_changed(project, ada)
    base = f"{server}/api/projects/{project}"
    tried = [(f"{base}/files/notes.txt", "PUT", b"x"), (f"{server}/projects/{project}/files/"
             "notes.txt", "GET", None), (f"{base}/history", "GET", None),
             (f"{base}/files/notes.txt", "DELETE", None)]  # fmt: skip
    assert [call(url, method, data)[0] for url, method, data in tried] == [404] * 4
    agent.change_project(project, ada, members=[1, ben])
    connection = http.client.HTTPConnection(server.replace("http://", ""), timeout=5)
    connection.request("GET", "/api/events", headers={"Cookie": cookie()})
    stream = connection.getresponse()
    while not stream.readline().startswith(b'data: {"type": "settings"'):  # all there was, told
        pass
    assert [call(url, method, data)[0] for url, method, data in tried] == [200] * 4
    told: list[dict] = []
    while not told or told[-1]["files"]:  # told of it uploaded, then deleted
        if (line := stream.readline()).startswith(b"data: "):
            told.append(json.loads(line[6:]))
    assert [e["files"][0]["name"] for e in told[:-1]] == ["notes.txt"]
    assert {e["project"] for e in told} == {project}
    connection.close()
    agent.change_project(project, ada, members=[ben])
    assert [call(url, method, data)[0] for url, method, data in tried] == [404] * 4
    with agent.events.watch() as events:
        agent._publish_files(project, None)
        assert events.get()["to"] == [ada, ben]


@sandboxed
def test_history(workspace, monkeypatch):
    # a file written over, edited or deleted kept first as an earlier version, apart from the
    # workspace; restored, the file as it was then kept too; a deleted one restored; hidden
    # files not kept; a space's history its own, and gone with it
    files.write(workspace, "notes.txt", "one")
    assert workspace.versions() == []  # a new file: nothing earlier
    files.write(workspace, "notes.txt", "two")
    files.edit(workspace, "notes.txt", "two", "three")
    earlier = workspace.versions()
    assert [(v["name"], v["size"], v["deleted"]) for v in earlier] == [
        ("notes.txt", 3, False), ("notes.txt", 3, False)]  # "two", then "one"  # fmt: skip
    assert not workspace.history.is_relative_to(workspace.root)
    workspace.restore("notes.txt", earlier[-1]["version"])
    assert (workspace.root / "notes.txt").read_text() == "one"
    assert len(workspace.versions()) == 3  # "three" kept as it was restored over
    (workspace.root / ".saved.md").write_text("a page read")  # hidden: deleted, not kept
    workspace.delete(".saved.md")
    assert not (workspace.root / ".saved.md").exists()
    assert {v["name"] for v in workspace.versions()} == {"notes.txt"}
    files.write(workspace, "old/a.txt", "a")
    (workspace.root / "old" / ".hidden").write_text("h")
    workspace.delete("old")
    gone = [v for v in workspace.versions() if v["deleted"]]
    assert [v["name"] for v in gone] == ["old/a.txt"] and not (workspace.root / "old").exists()
    workspace.restore("old/a.txt", gone[0]["version"])
    assert (workspace.root / "old" / "a.txt").read_text() == "a"
    with pytest.raises(FileNotFoundError):
        workspace.restore("notes.txt", "123")
    with pytest.raises(ValueError):
        workspace.restore("../x", "123")
    # a space's, apart from the others', and removed with it
    space = workspace.space("projects/a")
    files.write(space, "x.txt", "1")
    files.write(space, "x.txt", "2")
    assert space.history == workspace.history / "projects" / "a"
    assert [v["name"] for v in space.versions()] == ["x.txt"]
    space.remove()
    assert not space.history.exists() and not space.root.exists()
    # HISTORY versions of a file at most, and none kept longer than KEPT
    monkeypatch.setattr(workspace_module, "HISTORY", 2)
    files.write(workspace, "notes.txt", "four")
    assert len([v for v in workspace.versions() if v["name"] == "notes.txt"]) == 2
    monkeypatch.setattr(workspace_module, "KEPT", 0)
    workspace.forget()
    assert workspace.versions() == []


@sandboxed
def test_run_keeps(workspace):
    # code run changes the workspace only once it ends, through its overlay, each file it
    # changes or deletes kept first: a file written over, one deleted, a folder deleted and made
    # anew, a file made a folder, and a new file; the history out of its sight
    if not workspace_module._overlays(workspace.scratch):
        pytest.skip("the kernel mounts no overlay for an unprivileged user")
    for name, text in (("a.txt", "a"), ("b.txt", "b"), ("d/1.txt", "1"), ("d/2.txt", "2"),
                       ("e", "e")):  # fmt: skip
        files.write(workspace, name, text)
    code = """
import os, shutil
open("a.txt", "w").write("A")
os.remove("b.txt")
shutil.rmtree("d"); os.mkdir("d"); open("d/3.txt", "w").write("3")
os.remove("e"); os.mkdir("e"); open("e/f.txt", "w").write("f")
open("new.txt", "w").write("new")
print(os.path.exists("/workspace/../.workspace-history"))
"""
    ran = files.run(workspace, code)
    assert ran.info["status"] == 0 and ran.content.startswith("False\n")
    names = sorted(f["name"] for f in workspace.files())
    assert names == ["a.txt", "d/3.txt", "e/f.txt", "new.txt"]
    assert (workspace.root / "a.txt").read_text() == "A"
    kept = {v["name"]: v["deleted"] for v in workspace.versions()}
    assert kept == {"a.txt": False, "b.txt": True, "d/1.txt": True, "d/2.txt": True, "e": True}
    assert not list(workspace.scratch.iterdir())  # the run's overlay gone


def test_overview(agent):
    # the room the files take, and their earlier versions, and the disk's
    files.write(agent.space(), "a.txt", "12345")
    files.write(agent.space(), "a.txt", "1234567")  # its earlier version kept
    seen = agent.overview()
    assert (seen["files"], seen["versions"]) == (7, 5)
    assert seen["disk"]["free"] <= seen["disk"]["total"] and seen["people"] == []


def test_sandbox_sees_its_space(agent):
    # code a conversation runs sees its space's files, and no other's
    project = agent.add_project("Payroll")["id"]
    files.write(agent.space(), "mine.txt", "mine")
    files.write(agent.space(project), "theirs.txt", "theirs")
    ran = files.run(agent.space(project), "import os; print(sorted(os.listdir('.')))")
    assert ran.content.startswith("['theirs.txt']")


def test_folders(workspace):
    # an upload's folders kept, but for any that would lead out or are hidden, its name free in
    # the last; a folder deleted with its files
    assert workspace.free("Invoices/March/1.pdf", folders=True) == "Invoices/March/1.pdf"
    assert workspace.free("../.git/./a/1.pdf", folders=True) == "git/a/1.pdf"
    assert not (workspace.root / "Invoices").exists()  # named, not made
    files.write(workspace, "Invoices/March/1.pdf", "x")
    assert workspace.free("Invoices/March/1.pdf", folders=True) == "Invoices/March/1 (2).pdf"
    assert workspace.free("Invoices/March/1.pdf") == "1.pdf"
    workspace.delete("Invoices")
    assert workspace.files() == []
    with pytest.raises(FileNotFoundError):
        workspace.delete(".")


@sandboxed
def test_unpack(workspace, monkeypatch):
    # a zip unpacked into a folder of its name, its own folders kept, macOS's leavings, hidden
    # files and any name that would lead out left out; one too big, or broken, kept, and why said
    import zipfile

    with zipfile.ZipFile(workspace.root / "March.zip", "w") as z:
        z.writestr("Invoices/1.txt", "one")
        z.writestr("Fatura ğüş.txt", "two")
        z.writestr("__MACOSX/._1.txt", "junk")
        z.writestr("../../escape.txt", "out")
        z.writestr("Invoices/.DS_Store", "junk")
    assert workspace.unpack("March.zip") == ("March", 3)
    assert sorted(f["name"] for f in workspace.files()) == [
        "March/Fatura ğüş.txt", "March/Invoices/1.txt", "March/escape.txt"]  # fmt: skip
    (workspace.root / "broken.zip").write_bytes(b"not a zip")
    with pytest.raises(ValueError, match="broken.zip could not be unpacked"):
        workspace.unpack("broken.zip")
    monkeypatch.setattr(workspace_module, "MEMBERS", 1)
    with zipfile.ZipFile(workspace.root / "many.zip", "w") as z:
        z.writestr("a.txt", "a")
        z.writestr("b.txt", "b")
    with pytest.raises(ValueError, match="it holds 2 files, more than 1"):
        workspace.unpack("many.zip")
    assert (workspace.root / "many.zip").exists() and not (workspace.root / "many").exists()
    # a folder compressed, the folder that holds all the rest left out
    monkeypatch.setattr(workspace_module, "MEMBERS", 10)
    with zipfile.ZipFile(workspace.root / "Contracts.zip", "w") as z:
        z.writestr("Contracts/A.txt", "a")
        z.writestr("Contracts/2025/B.txt", "b")
    assert workspace.unpack("Contracts.zip") == ("Contracts", 2)
    assert sorted(f["name"] for f in workspace.files() if f["name"].startswith("Contracts/")) == [
        "Contracts/2025/B.txt", "Contracts/A.txt"]  # fmt: skip


@sandboxed
def test_api_folders(server, agent):
    # files uploaded in their folders, a zip unpacked if asked, a folder deleted
    import io
    import zipfile

    status, _, body = call(f"{server}/api/files/Invoices%2FMarch%2F1.txt", "PUT", b"one")
    assert status == 200 and json.loads(body) == {"name": "Invoices/March/1.txt"}
    packed = io.BytesIO()
    with zipfile.ZipFile(packed, "w") as z:
        z.writestr("2.txt", "two")
    status, _, body = call(f"{server}/api/files/April.zip?unpack", "PUT", packed.getvalue())
    assert status == 200 and json.loads(body) == {"name": "April", "files": 1}
    names = sorted(f["name"] for f in agent.space(None, 1).files())
    assert names == ["April/2.txt", "Invoices/March/1.txt"]
    assert call(f"{server}/api/files/Invoices", "DELETE")[0] == 200
    assert [f["name"] for f in agent.space(None, 1).files()] == ["April/2.txt"]
    status, _, body = call(f"{server}/api/files/bad.zip?unpack", "PUT", b"not a zip")
    assert status == 400 and b"could not be unpacked" in body


def test_activity(server, agent, engine, tmp_path):
    # what people do, and what Leat does for them, noted as it is done, never what was asked;
    # the admins' alone to see, the latest first, in pages
    from tests.test_agent import request

    agent.engine = Client(engine.url)
    agent.tools = {t.name: t for t in files.tools()}
    _, body = request(f"{server}/api/projects", "POST", {"name": "Yılmaz Ltd"})
    project = json.loads(body)["id"]
    call(f"{server}/api/projects/{project}/files/Lease.txt", "PUT", b"Rent: 900.")
    engine.replies.put([{"tool_calls": [calling("read", {"path": "Lease.txt"})]}])
    engine.replies.put([{"content": "Rent is 900."}])
    with agent.events.watch() as events:
        request(f"{server}/api/conversations", "POST", {"content": "Secret question about rent?",
                                                        "project": project})  # fmt: skip
        until(events, ended)
    request(f"{server}/api/projects/{project}", "PATCH", {"shared": True})
    request(f"{server}/api/settings", "PATCH", {"offline": True})
    status, body = request(f"{server}/api/activity")
    seen = [(a["person"], a["leat"], a["action"], a["place"], a["detail"])
            for a in json.loads(body)["activity"]]  # fmt: skip
    assert status == 200 and seen == [
        ("Alkın", 0, "turned the internet off", None, None),
        ("Alkın", 0, "shared the project with", "Yılmaz Ltd", "everyone"),
        ("Alkın", 1, "read", "Yılmaz Ltd", "Lease.txt"),
        ("Alkın", 0, "began a chat", "Yılmaz Ltd", None),
        ("Alkın", 0, "uploaded", "Yılmaz Ltd", "Lease.txt"),
        ("Alkın", 0, "made the project", "Yılmaz Ltd", None),
        ("Alkın", 0, "set Leat up", None, None),
    ]  # fmt: skip
    assert "Secret" not in body.decode() and "rent" not in body.decode().lower()
    newest = json.loads(body)["activity"][0]["id"]
    _, older = request(f"{server}/api/activity?before={newest}")
    assert json.loads(older)["activity"][0]["action"] == "shared the project with"
    assert request(f"{server}/api/activity?before=x")[0] == 400
    request(f"{server}/api/projects/{project}", "DELETE")
    _, body = request(f"{server}/api/activity")
    assert json.loads(body)["activity"][0]["action"] == "deleted the project"
    assert json.loads(body)["activity"][0]["place"] == "Yılmaz Ltd"


def test_missing(workspace):
    # a file there is not, said with the nearest names there are, as one named in another language
    files.write(workspace, "Kontoauszug März 2026.csv", "x")
    files.write(workspace, "Notes.txt", "x")
    with pytest.raises(FileNotFoundError, match="the nearest: Kontoauszug März 2026.csv$"):
        files.read(workspace, "Kontoauszug March 2026.csv")
    with pytest.raises(FileNotFoundError, match="^there is no file zzz$"):
        files.read(workspace, "zzz")
    files.write(workspace, "lease.docx", "not text")
    with pytest.raises(ValueError, match="suggest_edits suggests edits to a Word document"):
        files.edit(workspace, "lease.docx", "not", "now")

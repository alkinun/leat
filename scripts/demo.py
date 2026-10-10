"""Makes a demo of Leat for a small office of accountants and lawyers, in a new state of its own, in
English, of a firm in Manchester, or in German, of one in Munich: a client's books, its month's
invoices, two of them scanned, its bank statement and its books of the month, which differ where a
reconciliation should find them out, its lease and a fee letter's template; and an employment
case, its contract, with clauses an employee's lawyer would change. Each project has workflows of
what the office asks of it. Then `leat agent --data DIR` serves it, and whoever first opens the app
owns it, the demo theirs.

    uv run python scripts/demo.py DIR [--language en|de] [--sandbox ~/.local/share/leat/sandbox]

The documents are made in the sandbox, with its libraries, as the agent's are, in a font of the
box's. Asking the scanned invoices of their text needs a model that sees images.
"""

import argparse
import datetime
import shutil
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from leat.agent.agent import Agent  # noqa: E402
from leat.agent.client import Client  # noqa: E402
from leat.agent.store import Store  # noqa: E402
from leat.agent.workspace import Workspace  # noqa: E402

MONTH_END = datetime.date(2026, 3, 31)  # of the month the demo's books and bank are of
SHORT = 500.0  # paid less than the invoice that is paid short


def _march(day: int) -> datetime.date:
    # a day of the month the demo's books and bank are of
    return MONTH_END.replace(day=day)


@dataclass(frozen=True)
class Invoice:
    """An invoice of the month: its number, date, supplier, net amount and VAT in percent, and
    whether it is scanned, an image of its page rather than its text."""

    number: str
    date: datetime.date
    supplier: str
    net: float
    rate: int
    scanned: bool = False

    @property
    def vat(self) -> float:
        return round(self.net * self.rate / 100, 2)

    @property
    def gross(self) -> float:
        return round(self.net + self.vat, 2)


@dataclass(frozen=True)
class Line:
    """A line of the books and of the bank alike: its date, its reference, the books' account and
    what they say, what the bank says and of whom, and its amount, as the bank has it, less if
    paid out; and of the bank's, of what kind it is, as an English bank's export says."""

    date: datetime.date
    reference: str
    account: str
    text: str
    said: str
    party: str
    amount: float
    kind: str = "FPO"


class Office:
    """A demo's office, in its language: its client's books, and an employment case. Each
    language's office says what it holds, and how its documents write it."""

    client: str  # the client whose books the office keeps, its project's name
    books: str  # its project's instructions
    invoices: list[Invoice]
    # where the books differ from the bank, for the reconciliation to find: the invoice paid but not
    # booked, the one booked but never paid, and the one paid SHORT less than it was booked
    unbooked: str
    unpaid: str
    short: str
    others: list[Line]  # in the books and the bank alike, as the rent, the wages and the receipts
    charges: list[Line]  # in the bank alone, as its fees
    payable: str  # the account the books pay invoices from
    day: str  # a date's format
    folder: str  # the invoices'
    ledger: tuple[str, str, list[str]]  # the books' workbook's name, its sheet's, its heading
    statement: str  # the bank statement's name
    lease: tuple[str, list[tuple[str | None, str | None]]]  # a Word document's name and blocks
    letter: tuple[str, list[str]]  # the fee letter's template's name and paragraphs
    books_workflows: list[tuple[str, str]]  # each's name and request
    case: str  # the employment case, its project's name
    case_instructions: str
    contract: tuple[str, list[tuple[str | None, str | None]]]
    case_workflows: list[tuple[str, str]]

    def page(self, invoice: Invoice) -> list[str]:
        """An invoice's lines, as its page shows them."""
        raise NotImplementedError

    def paid(self, invoice: Invoice) -> str:
        """What the bank says of an invoice's payment."""
        raise NotImplementedError

    def bank(self, lines: list[Line]) -> tuple[str, str, list[list[str]]]:
        """The bank statement's encoding, its separator and rows, of its lines."""
        raise NotImplementedError

    def made(self) -> dict[str, dict[str, list]]:
        """The documents of each project, by its name, as the sandbox makes them."""
        name, sheet, heading = self.ledger
        books = [  # each in the books on the day it was paid, before it is in the bank
            [f"{line.date:{self.day}}", line.reference, line.account, line.text,
             line.amount if line.amount > 0 else None, -line.amount if line.amount < 0 else None]
            for line in sorted(self._lines(books=True), key=lambda line: line.date)
        ]  # fmt: skip
        return {
            self.client: {
                "pdfs": [
                    [f"{self.folder}/{i.number.replace('/', '-')}.pdf", self.page(i), i.scanned]
                    for i in self.invoices
                ],  # fmt: skip
                "workbooks": [[name, sheet, [heading, *books]]],
                "csvs": [[self.statement, *self.bank(self._lines(books=False))]],
                "documents": [
                    list(self.lease),
                    [self.letter[0], [(None, p) for p in self.letter[1]]],
                ],
            },
            self.case: {
                "pdfs": [],
                "workbooks": [],
                "csvs": [],
                "documents": [list(self.contract)],
            },
        }

    def _lines(self, books: bool) -> list[Line]:
        # the month's lines of the books, or of the bank: each invoice's but the one the other
        # lacks, its payment made three days after the invoice and in the bank a day later, in the
        # month; and the others, and the bank's charges
        lines = []
        for invoice in self.invoices:
            if invoice.number == (self.unbooked if books else self.unpaid):
                continue
            amount = invoice.gross - (SHORT if invoice.number == self.short and not books else 0)
            day = min(invoice.date + datetime.timedelta(days=3 if books else 4), MONTH_END)
            lines.append(Line(day, invoice.number, self.payable, invoice.supplier,
                              self.paid(invoice), invoice.supplier, -round(amount, 2),
                              "DD" if "Energy" in invoice.supplier else "FPO"))  # fmt: skip
        return lines + self.others + ([] if books else self.charges)


class English(Office):
    """A firm of accountants and solicitors in Manchester, its client a textile wholesaler."""

    client = "Hartley Textiles Ltd"
    books = """\
Our client is Hartley Textiles Ltd, a textile wholesaler in Manchester, company number 08412397, \
VAT number GB 284 7193 50; its financial year ends on 31 March. Answer in British English, write \
amounts as £1,234.56, name the document each figure comes from, and say where you are unsure."""
    invoices = [
        Invoice("INV-2026-0342", _march(2), "Northern Office Supplies Ltd", 1_000.00, 20),
        Invoice("PWC/118342", _march(3), "Pennine Weaving Co Ltd", 14_250.00, 20),
        Invoice("CF-77310", _march(5), "Calder Freight Ltd", 2_380.00, 20),
        Invoice("NG-5530182", _march(6), "Northgrid Energy Ltd", 1_424.50, 20),
        Invoice("KP-00932", _march(9), "Kowalski Packaging", 712.50, 20, scanned=True),
        Invoice("PWC/118690", _march(12), "Pennine Weaving Co Ltd", 9_800.00, 20),
        Invoice("DB-2026-215", _march(14), "Dyeworks Bradford Ltd", 3_960.00, 20),
        Invoice("INV-2026-0350", _march(18), "Mercer IT Services", 2_000.00, 20),
        Invoice("TM-44102", _march(20), "Trafford Machine Services", 315.00, 20,
                scanned=True),
        Invoice("CF-77391", _march(24), "Calder Freight Ltd", 1_130.00, 20),
        Invoice("0047", _march(26), "Ancoats Bakehouse", 448.00, 0),
        Invoice("NG-5531447", _march(30), "Northgrid Energy Ltd", 1_387.25, 20),
    ]  # fmt: skip
    unbooked, unpaid, short = "DB-2026-215", "TM-44102", "PWC/118690"
    others = [
        Line(_march(16), "H-2041", "Sales", "Harrogate Interiors Ltd H-2041",
             "HARROGATE INTERIOR H-2041", "", 18_600.00, "BGC"),
        Line(_march(20), "PAYE-02", "PAYE and NIC", "HMRC PAYE and NIC, February",
             "HMRC PAYE 120PA00123456", "", -9_880.00, "BP"),
        Line(_march(23), "B-1187", "Sales", "Bloom & Co B-1187",
             "BLOOM AND CO B-1187", "", 7_320.00, "BGC"),
        Line(_march(25), "RENT-Q2", "Rent", "Rent, Unit 4, 25 March to 23 June",
             "POMONA ESTATES RENT", "", -15_120.00),
        Line(_march(27), "PAYROLL-03", "Wages", "Net pay, March",
             "PAYROLL MARCH 2026", "", -31_250.00, "BP"),
    ]  # fmt: skip
    charges = [Line(MONTH_END, "", "", "", "ACCOUNT FEE MARCH", "", -12.50, "FEE")]
    payable, day, folder = "Accounts payable", "%d/%m/%Y", "Invoices"
    ledger = ("Cash book March 2026.xlsx", "Cash book",
              ["Date", "Reference", "Account", "Description", "Debit", "Credit"])  # fmt: skip
    statement = "Bank statement March 2026.csv"
    lease = ("Warehouse lease.docx", [
        ("Lease of Unit 4, Pomona Business Park", None),
        ("1. Parties", "This lease is made between Pomona Estates Ltd (the Landlord) and Hartley "
         "Textiles Ltd (the Tenant)."),
        ("2. Premises", "Unit 4, Pomona Business Park, Trafford Road, Salford M5 3AW: a warehouse "
         "of 1,850 square metres with its loading yard."),
        ("3. Term", "Ten years from 25 March 2026."),
        ("4. Rent", "£50,400 a year plus VAT, payable quarterly in advance on the usual quarter "
         "days."),
        ("5. Deposit", "The Tenant shall pay the Landlord a deposit equal to six months' rent."),
        ("6. Break clause", "The Tenant may end this lease on the fifth anniversary of the term by "
         "giving the Landlord at least twelve months' written notice, provided that it has paid "
         "all rent due."),
        ("7. Repair", "The Tenant shall keep the premises, including the roof and structure, in "
         "good and substantial repair."),
        ("8. Rent review", "The rent shall be reviewed on the fifth anniversary of the term to the "
         "open market rent, upwards only."),
        ("9. Governing law", "This lease is governed by the law of England and Wales."),
    ])  # fmt: skip
    letter = ("Fee letter template.docx", [
        "Hartley Textiles Ltd",
        "For the attention of {{Contact}}",
        "Fees for {{Month}}",
        "Dear {{Contact}},",
        "Our fee for bookkeeping, VAT and payroll for {{Month}} is {{Fee}} plus VAT, payable by "
        "{{Due date}}.",
        "Yours sincerely,",
        "{{Partner}}",
    ])  # fmt: skip
    books_workflows = [
        ("Invoice summary", "Make a table of every invoice in the Invoices folder with its "
         "number, date, supplier, VAT and total, and name the three suppliers we paid the most."),
        ("Bank reconciliation", "Reconcile the March 2026 bank statement with the March 2026 cash "
         "book and explain each difference."),
        ("Fee letter", "Fill in the fee letter template for March 2026: fee £1,850, due 10 April "
         "2026, for the attention of Emma Hartley, signed by Sarah Whitfield."),
    ]  # fmt: skip
    case = "Nowak v Bramley Logistics Ltd"
    case_instructions = """\
We act for the employee, Tomasz Nowak, against his employer, Bramley Logistics Ltd. Review clauses \
under the employment law of England and Wales, name the statute or rule you rely on, and answer in \
English unless asked otherwise."""
    contract = ("Employment contract.docx", [
        ("Contract of Employment", None),
        ("1. Parties", "This contract is between Bramley Logistics Ltd, Leeds (the Company), and "
         "Tomasz Nowak (the Employee). The Employee's employment began on 3 February 2020."),
        ("2. Role", "The Employee is employed as a Transport Planner. The Company may at any time "
         "require the Employee to carry out any other duties at any of its sites in the United "
         "Kingdom."),
        ("3. Hours", "The Employee's normal hours are 45 a week. No additional payment is made for "
         "overtime."),
        ("4. Pay", "The Employee is paid £29,500 a year, monthly in arrears."),
        ("5. Holiday", "The Employee is entitled to 18 days' paid holiday a year, including bank "
         "holidays."),
        ("6. Notice", "Either party may end this contract by giving one week's written notice, "
         "whatever the Employee's length of service."),
        ("7. Restrictions", "For 24 months after this contract ends, the Employee shall not work "
         "for any business in the logistics sector anywhere in the United Kingdom."),
        ("8. Deductions", "The Company may deduct from the Employee's pay the full cost of any "
         "damage to its vehicles, however caused."),
        ("9. Claims", "Any claim arising from this contract must be raised in writing within four "
         "weeks, after which it is waived."),
    ])  # fmt: skip
    case_workflows = [
        ("Contract review", "Review the employment contract for clauses that are unenforceable or "
         "unfavourable to our client, and suggest changes as tracked changes in Word, each with "
         "its reason and the statute or rule it rests on."),
        ("Translation", "Translate the employment contract into Polish for our client."),
    ]  # fmt: skip

    def page(self, invoice: Invoice) -> list[str]:
        vat = (f"VAT at {invoice.rate}%: {_pounds(invoice.vat)}" if invoice.rate
               else "No VAT charged: the supplier is not VAT registered")  # fmt: skip
        return [
            f"INVOICE  {invoice.number}", f"Invoice date: {invoice.date:%d/%m/%Y}",
            f"From: {invoice.supplier}", "To: Hartley Textiles Ltd, 14 Jersey Street, "
            "Manchester M4 6JG", "", f"Net amount: {_pounds(invoice.net)}", vat,
            f"Total due: {_pounds(invoice.gross)}", "", "Payment due within 14 days.",
        ]  # fmt: skip

    def paid(self, invoice: Invoice) -> str:
        # as a British bank cuts a payee's name
        return f"{invoice.supplier.upper()[:18].strip()} {invoice.number}"

    def bank(self, lines: list[Line]) -> tuple[str, str, list[list[str]]]:
        # as a British bank's export has them, the latest first, each with the balance after it
        rows, balance = [], 120_000.00
        for line in sorted(lines, key=lambda line: line.date):
            balance = round(balance + line.amount, 2)
            amount = f"{abs(line.amount):.2f}"
            out, into = (amount, "") if line.amount < 0 else ("", amount)
            rows.append([f"{line.date:%d/%m/%Y}", line.kind, "30-94-57", "12345678", line.said, out,
                         into, f"{balance:.2f}"])  # fmt: skip
        heading = [
            "Transaction Date", "Transaction Type", "Sort Code", "Account Number",
            "Transaction Description", "Debit Amount", "Credit Amount", "Balance",
        ]  # fmt: skip
        return "utf-8", ",", [heading, *reversed(rows)]


class German(Office):
    """A law and tax office in Munich, its client a textile wholesaler."""

    client = "Hofmann Textil GmbH"
    books = """\
Mandantin ist die Hofmann Textil GmbH, ein Textilgroßhandel in München, USt-IdNr. DE284719350; \
das Geschäftsjahr endet am 30. Juni. Antworte auf Deutsch, schreibe Beträge wie 1.234,56 EUR, \
nenne zu jeder Zahl den Beleg, aus dem sie stammt, und sage, wo du unsicher bist."""
    invoices = [
        Invoice("RE-2026-0342", _march(2), "Müller Bürobedarf GmbH", 1_000.00, 19),
        Invoice("2026-118342", _march(3), "Weberei Lindner KG", 14_250.00, 19),
        Invoice("LG-77310", _march(5), "Weber Logistik KG", 2_380.00, 19),
        Invoice("SWM-5530182", _march(6), "Stadtwerke München", 1_424.50, 19),
        Invoice("KA-00932", _march(9), "Kaya Verpackungen", 712.50, 19, scanned=True),
        Invoice("2026-118690", _march(12), "Weberei Lindner KG", 9_800.00, 19),
        Invoice("FB-2026-215", _march(14), "Färberei Brandt GmbH", 3_960.00, 19),
        Invoice("RE-2026-0350", _march(18), "Schmidt IT-Service", 2_000.00, 19),
        Invoice("TS-44102", _march(20), "Technik Service Ost", 315.00, 19,
                scanned=True),
        Invoice("LG-77391", _march(24), "Weber Logistik KG", 1_130.00, 19),
        Invoice("BK-7781", _march(26), "Bäckerei Kraus", 448.00, 7),
        Invoice("SWM-5531447", _march(30), "Stadtwerke München", 1_387.25, 19),
    ]  # fmt: skip
    unbooked, unpaid, short = "FB-2026-215", "TS-44102", "2026-118690"
    others = [
        Line(_march(3), "MIETE-03", "4210 Miete", "Miete März Lagerhalle",
             "Miete März", "Kühn Immobilien GmbH", -4_998.00),
        Line(_march(10), "LSt-02", "1741 Lohnsteuer", "Lohnsteuer Februar",
             "Lohnsteuer 02/2026 StNr 143/123/45678", "Finanzamt München", -9_880.00),
        Line(_march(16), "AR-2026-041", "1400 Forderungen",
             "Modehaus Albrecht AR-2026-041", "AR-2026-041 Zahlung", "Modehaus Albrecht GmbH",
             18_600.00),
        Line(_march(23), "AR-2026-044", "1400 Forderungen",
             "Textilhaus Berger AR-2026-044", "AR-2026-044", "Textilhaus Berger KG", 7_320.00),
        Line(_march(31), "LOHN-03", "1740 Löhne", "Löhne März", "Löhne März",
             "Sammelüberweisung", -38_640.00),
    ]  # fmt: skip
    charges = [Line(MONTH_END, "", "", "", "Entgelt Kontoführung 03/2026", "Stadtsparkasse München",
                    -12.90)]  # fmt: skip
    payable, day, folder = "1600 Verbindlichkeiten", "%d.%m.%Y", "Rechnungen"
    ledger = ("Buchungen März 2026.xlsx", "Buchungen",
              ["Belegdatum", "Belegnr.", "Konto", "Buchungstext", "Soll", "Haben"])  # fmt: skip
    statement = "Kontoauszug März 2026.csv"
    lease = ("Gewerbemietvertrag.docx", [
        ("Gewerbemietvertrag", None),
        ("§ 1 Mietsache", "Vermietet wird die Lagerhalle Am Gewerbering 12, 85748 Garching, von "
         "der Kühn Immobilien GmbH an die Hofmann Textil GmbH."),
        ("§ 2 Mietzeit", "Das Mietverhältnis beginnt am 1. März 2026 und läuft auf unbestimmte "
         "Zeit."),
        ("§ 3 Miete", "Die monatliche Miete beträgt 4.200 EUR zuzüglich Umsatzsteuer und ist bis "
         "zum dritten Werktag eines jeden Monats im Voraus zu zahlen."),
        ("§ 4 Kaution", "Die Mieterin leistet eine Kaution in Höhe von drei Monatsmieten."),
        ("§ 5 Kündigung", "Die Kündigungsfrist beträgt sechs Monate zum Ende eines "
         "Kalendervierteljahres."),
        ("§ 6 Gerichtsstand", "Gerichtsstand ist München."),
    ])  # fmt: skip
    letter = ("Honorarschreiben Vorlage.docx", [
        "Hofmann Textil GmbH",
        "z. Hd. {{Ansprechpartner}}",
        "Honorarabrechnung {{Monat}}",
        "Sehr geehrte Frau Hofmann,",
        "für die laufende Finanzbuchhaltung und Lohnabrechnung im {{Monat}} berechnen wir Ihnen "
        "{{Honorar}} zuzüglich Umsatzsteuer, zahlbar bis zum {{Fälligkeit}}.",
        "Mit freundlichen Grüßen",
        "{{Partner}}",
    ])  # fmt: skip
    books_workflows = [
        ("Rechnungsübersicht", "Erstelle eine Tabelle aller Rechnungen im Ordner Rechnungen mit "
         "Rechnungsnummer, Datum, Lieferant, Umsatzsteuer und Bruttobetrag, und nenne die drei "
         "Lieferanten, an die wir am meisten gezahlt haben."),
        ("Kontoabstimmung", "Stimme den Kontoauszug März 2026 mit den Buchungen März 2026 ab und "
         "erkläre jede Abweichung."),
        ("Honorarschreiben", "Fülle die Vorlage für das Honorarschreiben für März 2026 aus: "
         "Honorar 1.850 EUR, fällig am 10. April 2026, Ansprechpartnerin Anna Hofmann, gezeichnet "
         "von Dr. Clara Becker."),
    ]  # fmt: skip
    case = "Schulz ./. Bauer Logistik GmbH"
    case_instructions = """\
Wir vertreten den Arbeitnehmer Markus Schulz gegen seine Arbeitgeberin, die Bauer Logistik GmbH. \
Prüfe Klauseln nach deutschem Arbeitsrecht, nenne die Vorschrift, auf die du dich stützt, und \
antworte auf Deutsch, außer man bittet dich um eine andere Sprache."""
    contract = ("Arbeitsvertrag.docx", [
        ("Arbeitsvertrag", None),
        ("§ 1 Parteien", "Zwischen der Bauer Logistik GmbH, Hamburg (Arbeitgeberin), und Herrn "
         "Markus Schulz (Arbeitnehmer) wird folgender Arbeitsvertrag geschlossen."),
        ("§ 2 Tätigkeit", "Der Arbeitnehmer wird als Disponent eingestellt. Die Arbeitgeberin kann "
         "ihm jederzeit jede andere Tätigkeit an jedem Ort zuweisen."),
        ("§ 3 Arbeitszeit", "Die regelmäßige Arbeitszeit beträgt 40 Stunden in der Woche. "
         "Überstunden sind mit dem Gehalt abgegolten."),
        ("§ 4 Vergütung", "Der Arbeitnehmer erhält ein Bruttomonatsgehalt von 3.900 EUR."),
        ("§ 5 Urlaub", "Der Arbeitnehmer hat Anspruch auf 18 Arbeitstage Urlaub im Kalenderjahr."),
        ("§ 6 Kündigung", "Das Arbeitsverhältnis kann von beiden Seiten mit einer Frist von zwei "
         "Wochen gekündigt werden."),
        ("§ 7 Wettbewerbsverbot", "Der Arbeitnehmer darf zwei Jahre nach Ende des "
         "Arbeitsverhältnisses für kein Unternehmen der Logistikbranche tätig werden."),
        ("§ 8 Ausschlussfrist", "Ansprüche aus dem Arbeitsverhältnis verfallen, wenn sie nicht "
         "innerhalb von vier Wochen nach Fälligkeit schriftlich geltend gemacht werden."),
    ])  # fmt: skip
    case_workflows = [
        ("Vertragsprüfung", "Prüfe den Arbeitsvertrag auf Klauseln, die unwirksam oder für unseren "
         "Mandanten nachteilig sind, und schlage Änderungen als Nachverfolgung in Word vor, jede "
         "mit einer Begründung und der Vorschrift."),
        ("Übersetzung", "Übersetze den Arbeitsvertrag ins Englische, für den Konzernanwalt der "
         "Arbeitgeberin."),
    ]  # fmt: skip

    def page(self, invoice: Invoice) -> list[str]:
        return [
            f"RECHNUNG  {invoice.number}", f"Rechnungsdatum: {invoice.date:%d.%m.%Y}",
            f"Lieferant: {invoice.supplier}",
            "Rechnungsempfänger: Hofmann Textil GmbH, Lindwurmstraße 88, 80337 München", "",
            f"Nettobetrag: {_euros(invoice.net)}",
            f"Umsatzsteuer {invoice.rate} %: {_euros(invoice.vat)}",
            f"Rechnungsbetrag: {_euros(invoice.gross)}", "",
            "Zahlbar innerhalb von 14 Tagen ohne Abzug.",
        ]  # fmt: skip

    def paid(self, invoice: Invoice) -> str:
        return f"{invoice.number} {invoice.supplier}"

    def bank(self, lines: list[Line]) -> tuple[str, str, list[list[str]]]:
        # as a Sparkasse's export has them, of Windows' Western European
        heading = ["Auftragskonto", "Buchungstag", "Verwendungszweck",
                   "Begünstigter/Zahlungspflichtiger", "Betrag", "Währung"]  # fmt: skip
        rows = [["DE89701500000012345678", f"{line.date:%d.%m.%Y}", line.said, line.party,
                 _euros(line.amount, unit=False), "EUR"]
                for line in sorted(lines, key=lambda line: line.date)]  # fmt: skip
        return "cp1252", ";", [heading, *rows]


OFFICES = {"en": English, "de": German}
# makes the documents of the JSON at sys.argv[1], in the font it names, in the sandbox: PDFs,
# of lines, each of its text or, scanned, an image of it; workbooks, of a sheet's rows, the first
# its heading; CSVs, of rows, in an encoding, of a separator; and Word documents, of headings and
# paragraphs
_MAKE = """
import csv, json, sys
from pathlib import Path
from docx import Document
from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFont

spec = json.load(open(sys.argv[1], encoding="utf-8"))
font = spec["font"]
for name, lines, scanned in spec["pdfs"]:
    Path(name).parent.mkdir(parents=True, exist_ok=True)
    pdf = FPDF()
    pdf.add_font("Sans", fname=font)
    pdf.add_page()
    if scanned:  # an image of the page, slightly askew, as a scanner's
        page = Image.new("L", (1240, 1754), 248)
        draw, face = ImageDraw.Draw(page), ImageFont.truetype(font, 34)
        for i, line in enumerate(lines):
            draw.text((110, 140 + i * 62), line, fill=30, font=face)
        page.rotate(0.6, fillcolor=248).save(".scan.png")
        pdf.image(".scan.png", x=0, y=0, w=210)
        Path(".scan.png").unlink()
    else:
        pdf.set_font("Sans", size=16)
        pdf.cell(text=lines[0], new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Sans", size=11)
        for line in lines[1:]:
            if line:
                pdf.cell(text=line, new_x="LMARGIN", new_y="NEXT", h=8)
            else:
                pdf.ln(8)
    pdf.output(name)
for name, title, rows in spec["workbooks"]:
    book = Workbook()
    sheet = book.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(50, width + 2)
    book.save(name)
for name, encoding, separator, rows in spec["csvs"]:
    with open(name, "w", encoding=encoding, newline="") as f:
        csv.writer(f, delimiter=separator, lineterminator="\\r\\n").writerows(rows)
for name, blocks in spec["documents"]:
    document = Document()
    for heading, text in blocks:
        if heading is not None:
            document.add_heading(heading, level=0 if text is None else 1)
        if text is not None:
            document.add_paragraph(text)
    document.save(name)
print("{}")
"""


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("data", type=Path, help="a new folder for the demo's state")
    parser.add_argument(
        "--language", choices=OFFICES, default="en",
        help="en, a firm in Manchester, by default; or de, an office in Munich",
    )  # fmt: skip
    parser.add_argument(
        "--sandbox", type=Path, default=Path.home() / ".local/share/leat/sandbox",
        help="the sandbox's environment, as leat/agent/sandbox.txt makes it",
    )  # fmt: skip
    args = parser.parse_args()
    if args.data.exists() and any(args.data.iterdir()):
        raise SystemExit(f"{args.data} is not empty: give a new folder for the demo")
    if not args.sandbox.exists():
        raise SystemExit(f"there is no sandbox environment at {args.sandbox}: make it as the "
                         "README says")  # fmt: skip
    office, font = OFFICES[args.language](), _font()
    args.data.mkdir(parents=True, exist_ok=True)
    (args.data / "sandbox").symlink_to(args.sandbox.resolve())  # leat agent's, of --data
    workspace = Workspace(args.data / "workspace", args.sandbox)
    agent = Agent(Store(args.data / "leat.db"), Client("http://127.0.0.1:9"), [], workspace)
    projects = {
        office.client: (office.books, office.books_workflows),
        office.case: (office.case_instructions, office.case_workflows),
    }
    made = office.made()
    try:
        for name, (instructions, workflows) in projects.items():
            project = agent.add_project(name, None, instructions, shared=True)["id"]
            agent.space(project).given(_MAKE, made[name] | {"font": font}, timeout=300)
            for workflow, prompt in workflows:
                agent.add_workflow(workflow, prompt, None, project)
            names = sorted(f["name"] for f in agent.space(project).files())
            print(f"{name}, {len(names)} {'file' if len(names) == 1 else 'files'}:")
            print("\n".join(f"  {name}" for name in names))
    except ValueError as e:
        shutil.rmtree(args.data)  # rather than a demo half made
        raise SystemExit(f"the documents could not be made: {e}") from e
    print(f"\nServe it with: uv run leat agent --data {args.data}")


def _font() -> str:
    # a font of the box's with German's letters and the pound's sign, its file's path, under /usr,
    # which the sandbox sees
    found = subprocess.run(["fc-match", "-f", "%{file}", "sans-serif:lang=de"],
                           capture_output=True, text=True, check=False).stdout  # fmt: skip
    if not found.startswith("/usr/") or not Path(found).is_file():
        raise SystemExit("no font was found: install one, as Noto Sans or DejaVu Sans")
    return found


def _pounds(amount: float) -> str:
    # an amount as British English writes it, £1,234.56
    return f"£{amount:,.2f}"


def _euros(amount: float, unit: bool = True) -> str:
    # an amount as German writes it, 1.234,56, with its currency if `unit`
    said = f"{amount:,.2f}".replace(",", "x").replace(".", ",").replace("x", ".")
    return f"{said} EUR" if unit else said


if __name__ == "__main__":
    main()

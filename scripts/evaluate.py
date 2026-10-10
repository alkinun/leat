"""Measures how the agent does what people ask of it: whether it calls the tools it should, and
says and makes what it should, each case run several times, in a state of its own. Prints a
Markdown table of each case's passes, mean time, and the share of the prompts' tokens the engine's
cache held.

    uv run python scripts/evaluate.py [--engine http://127.0.0.1:8080] [-n 3] [--effort E] [-k name]

The agent runs in this process, with leat agent's tools, against leat serve at --engine, the SearXNG
at --search, and the sandbox's environment at --sandbox, with leat serve's key in LEAT_ENGINE_KEY if
it asks for one, as leat agent's; the states are temporary, and the user's own untouched. A
prompt's or a tool's change is measured here before it is kept.
"""

import argparse
import os
import re
import statistics
import sys
import tempfile
import time
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from leat.agent.agent import Agent  # noqa: E402
from leat.agent.client import Client  # noqa: E402
from leat.agent.index import Index  # noqa: E402
from leat.agent.store import Store  # noqa: E402
from leat.agent.tools import ask, files, office, web  # noqa: E402
from leat.agent.workspace import Workspace  # noqa: E402
from leat.chat import EFFORTS  # noqa: E402

TIMEOUT = 300  # seconds a turn may take


@dataclass(frozen=True)
class Outcome:
    """What came of a case's message: the tools it called, the answer, and the files it made."""

    tools: list[str]
    answer: str
    files: list[str]
    seconds: float
    cached: float  # of the prompts' tokens, the share the engine's cache held


Check = Callable[[Outcome], str | None]  # why an outcome fails, or None if it passes


def called(name: str, times: int = 1) -> Check:
    n = lambda o: o.tools.count(name)  # noqa: E731
    return lambda o: None if n(o) >= times else f"called {name} {n(o)} times, not {times}"


def uncalled(*names: str) -> Check:
    return lambda o: next((f"called {name}" for name in names if name in o.tools), None)


def says(pattern: str) -> Check:
    return lambda o: None if re.search(pattern, o.answer, re.I) else f"did not say /{pattern}/"


def makes(pattern: str) -> Check:
    made = lambda o: any(re.search(pattern, name) for name in o.files)  # noqa: E731
    return lambda o: None if made(o) else f"made no file /{pattern}/"


def unsure(o: Outcome) -> str | None:
    # an answer that says it does not know, rather than one made up
    known = r"(n['’]t|not)( \w+){0,2} (know|have|remember|told|mention)|not sure"
    return None if re.search(known, o.answer, re.I) else "did not say it does not know"


@dataclass(frozen=True)
class Case:
    name: str
    message: str  # in a new conversation
    checks: list[Check]
    files: dict[str, str] = field(default_factory=dict)  # the workspace's, by name
    attach: bool = True  # whether the message attaches them, or the model must find them
    setup: str = ""  # Python run in the sandbox, in the workspace, that makes more of its files


# a lease, as a Word document, made in the sandbox, its Turkish and its English
LEASE = """
from docx import Document
d = Document()
d.add_heading("Kira Sözleşmesi", level=1)
d.add_paragraph("Kiracı: Yılmaz Tekstil Ltd. Kiraya veren: Bursa Gayrimenkul A.Ş.")
d.add_paragraph("Aylık kira 40.000 TL olup her ayın 5'inde ödenir.")
d.add_paragraph("Depozito iki aylık kira tutarındadır.")
d.add_paragraph("Taraflar doksan gün önceden yazılı bildirimle sözleşmeyi feshedebilir.")
d.save("Kira Sözleşmesi.docx")
d = Document()
d.add_paragraph("Dear {{Client name}},")
d.add_paragraph("Your fee for {{Month}} is {{Fee}}, due on {{Due date}}.")
d.add_paragraph("Kind regards, {{Partner}}")
d.save("Fee letter template.docx")
"""
# a firm's files, among which one passage answers each question
FIRM = {
    "Policies/Expenses.md": "# Expenses\n\nTravel is reimbursed at 8 TL a kilometre. Meals on "
    "client visits are reimbursed up to 750 TL a day, with a receipt.",
    "Policies/Leave.md": "# Annual leave\n\nStaff have 20 days of paid leave a year, and 5 more "
    "after five years. Leave is asked for two weeks ahead.",
    "Clients/Yılmaz Tekstil.md": "# Yılmaz Tekstil\n\nContact: Ayşe Yılmaz. VAT number "
    "8340021957. Their year ends in June; their accounts are due by 30 September.",
    "Clients/Ege Lojistik.md": "# Ege Lojistik\n\nContact: Mehmet Demir. Payroll for 42 staff, "
    "run on the 25th.",
}
# invoices, a file each
INVOICES = {
    f"Invoices/{n}.txt": f"FATURA {n}\nTarih: {date}\nSatıcı: {seller}\nToplam: {total} TL"
    for n, date, seller, total in [
        ("2026-031", "02.03.2026", "Akın Tekstil", "12.400,00"),
        ("2026-032", "05.03.2026", "Bursa Kumaş", "8.150,50"),
        ("2026-033", "09.03.2026", "Ege Lojistik", "2.300,00"),
        ("2026-034", "14.03.2026", "Akın Tekstil", "15.020,00"),
        ("2026-035", "21.03.2026", "Marmara Enerji", "4.870,25"),
    ]
}


NONE = ("search", "fetch", "read", "run")
CASES = [
    Case("chat", "Write a haiku about autumn.", [uncalled(*NONE)]),
    Case("arithmetic", "What is 17 * 23?", [says(r"391"), uncalled(*NONE)]),
    Case("known fact", "What is the capital of Australia?", [says("Canberra")]),
    Case("news", "What's in the news today about space exploration?", [called("search")]),
    Case("reads pages", "When does the British Museum open tomorrow? Check its website.",
         [called("fetch")]),
    Case("research", "Research the pros and cons of heat pumps for a house in a cold climate.",
         [called("search", 2), called("fetch", 4), says(r"\[\d+\]")]),
    Case("attachment", "What time does it start, and what should I bring?",
         [called("read"), says("7"), says("salad|dessert")],
         files={"invitation.txt": "You're invited to Mia's 30th! Saturday 18 October, 7 pm, at "
                "22 Oak Road. Please bring a salad or a dessert."}),
    Case("document", "Make a Word document with a packing list for a weekend camping trip.",
         [called("run"), makes(r"\.docx$")]),
    Case("spreadsheet", "Make an Excel budget: rent 900, food 350 and transport 80 a month, with "
         "yearly totals.", [called("run"), makes(r"\.xlsx$")]),
    Case("doesn't know", "What's my sister's name?", [unsure, uncalled("search")]),
    Case("finds in files", "How much are meals on client visits reimbursed?",
         [called("search_files"), says("750"), says(r"\[\d+\]")], FIRM, attach=False),
    Case("finds in Turkish", "Yılmaz Tekstil'in hesapları ne zamana kadar teslim edilmeli?",
         [called("search_files"), says("30 Eylül|30 September|30\\.09")], FIRM, attach=False),
    Case("not in files", "What is our policy on working from home?",
         [called("search_files"), unsure], FIRM, attach=False),
    Case("every file", "Make a table of every invoice's date, seller and total.",
         [called("ask_files"), makes(r"\.xlsx$"), says("15.020|15,020")], INVOICES, attach=False),
    Case("fills a template", "Fill the fee letter template for Ege Lojistik: March, 6.000 TL, due "
         "10 April, signed by Ayşe Kaya.", [called("fill_template"), makes(r"Ege.*\.docx$")],
         setup=LEASE, attach=False),
    Case("redlines", "Suggest changes to the lease: thirty days' notice instead of ninety, and a "
         "deposit of one month.", [called("suggest_edits"), makes(r"\.docx$")], setup=LEASE,
         attach=False),
    Case("translates", "Translate the lease into English.",
         [called("translate_document"), makes(r"\.docx$")], setup=LEASE,
         attach=False),
]  # fmt: skip


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--engine", default="http://127.0.0.1:8080", help="leat serve's address")
    parser.add_argument("--search", default="http://127.0.0.1:8888", help="a SearXNG's address")
    parser.add_argument(
        "--sandbox", type=Path, default=Path.home() / ".local/share/leat/sandbox",
        help="the sandbox's environment, as leat/agent/sandbox.txt makes it",
    )  # fmt: skip
    parser.add_argument("-n", "--runs", type=int, default=3, help="runs of each case")
    parser.add_argument(
        "--effort", choices=EFFORTS, help="the effort the model reasons at; by default its own"
    )
    parser.add_argument("-k", help="the cases whose names hold this alone")
    args = parser.parse_args()
    cases = [case for case in CASES if not args.k or args.k in case.name]
    rows, passed = [], 0
    for case in cases:
        failures, seconds, cached = [], [], []
        for _ in range(args.runs):
            outcome = _run(case, args)
            seconds.append(outcome.seconds)
            cached.append(outcome.cached)
            why = [w for check in case.checks if (w := check(outcome))]
            failures += why[:1]
            print(f"{case.name}: {why[0] if why else 'passed'}", file=sys.stderr, flush=True)
        passed += args.runs - len(failures)
        failed = "; ".join(sorted(set(failures)))
        rows.append(f"| {case.name} | {args.runs - len(failures)}/{args.runs} | "
                    f"{statistics.mean(seconds):.1f} | {statistics.mean(cached):.0%} | "
                    f"{failed} |")  # fmt: skip
    mode = f"reasoning at {args.effort}" if args.effort else "reasoning at the model's default"
    print(f"\n{passed} of {len(cases) * args.runs} passed, {mode}\n")
    print("| case | passed | mean s | cached | failures |\n|---|---:|---:|---:|---|")
    print("\n".join(rows))


def _run(case: Case, args: argparse.Namespace) -> Outcome:
    # the case, in an agent of a state of its own
    with tempfile.TemporaryDirectory() as data:
        environment = args.sandbox if args.sandbox.exists() else None
        workspace = Workspace(Path(data) / "workspace", environment)
        engine = Client(args.engine, os.environ.get("LEAT_ENGINE_KEY"))
        index = Index(Path(data) / "index.db", workspace)
        tools = [
            *web.tools(args.search, engine), *files.tools(index), *ask.tools(engine, index),
            *office.tools(engine),
        ]  # fmt: skip
        agent = Agent(Store(Path(data) / "leat.db"), engine, tools, workspace, index)
        space = agent.space()  # no one's own, as the conversation is
        for name, text in case.files.items():
            space.path(name).parent.mkdir(parents=True, exist_ok=True)
            space.path(name).write_text(text)
        if case.setup and space.run(case.setup).status != 0:
            raise RuntimeError(f"{case.name}'s setup failed")
        index.update("people/0")  # read before the model searches them
        before = {f["name"] for f in space.files()}
        attached = list(case.files) if case.attach else []
        start = time.monotonic()
        messages = _wait(agent, agent.send(None, case.message, args.effort, attached))
        seconds = time.monotonic() - start
        tools_called = [m["name"] for m in messages if m["role"] == "tool"]
        answer = messages[-1]["content"] if messages and messages[-1]["role"] == "assistant" else ""
        names = [f["name"] for f in space.files() if f["name"] not in before]  # made by it
        infos = [m["info"] for m in messages if m["role"] == "assistant" and "read" in m["info"]]
        held, read = sum(i["cached"] or 0 for i in infos), sum(i["read"] for i in infos)
        share = held / (held + read) if held + read else 0.0
        return Outcome(tools_called, answer or "", names, seconds, share)


def _wait(agent: Agent, id: str) -> list[dict]:
    # a conversation's messages once its turn has ended, or none if it failed and took it back
    end = time.monotonic() + TIMEOUT
    while (c := agent.conversation(id)) is not None and c["running"]:
        if time.monotonic() > end:
            agent.stop(id)
        time.sleep(0.1)
    return c["messages"] if c else []


if __name__ == "__main__":
    main()

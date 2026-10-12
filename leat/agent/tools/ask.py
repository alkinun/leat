"""Asking every file the same: ask_files has a reader, the model in a conversation of its own, read
each of the conversation's files, or those the call names, and answer for it, as leat.agent.tools.
web's reader answers for a page. The answers make a table, a row a file, each citing its file,
which the call returns and saves as a spreadsheet. A question of every file, as which contracts
renew themselves, or each invoice's date and total, is so answered whole, where a search finds
some of the files' passages alone.

The readers run READERS at once, which the engine batches, and the call tells how far it is as
each ends; once the turn is stopped, it reads no more. A file longer than a reader reads is read as
its start and the passages of it the index finds likeliest to answer.

The call says which files answered nothing, and which could not be read or asked, so that the model
passes over none: leaving out what the files say is how agents fail most at an office's work. A
column of amounts, each with a currency or cents, is totalled by code, read as each file writes it,
1.234,56 or 1,234.56, as a model adds up many figures less surely; the spreadsheet holds them as
numbers, to sum and sort.
"""

import concurrent.futures
import json
import mimetypes
import re
from typing import Any

from leat.agent import documents
from leat.agent.client import BACKGROUND, Client, EngineError
from leat.agent.context import data_url, picture
from leat.agent.index import Index
from leat.agent.tools import Context, Result, Tool
from leat.agent.workspace import Workspace

READERS = 4  # files read at once, as many as leat serve's slots by default
FILES = 500  # files a call asks at most
READ = 24000  # characters of a file a reader reads at most
START = 4000  # of a longer one's, its start, before the passages likeliest to answer
ANSWER = 500  # tokens of a reader's answer at most
SHOWN = 50  # rows of the table the model reads; the spreadsheet holds every one
CELL = 300  # characters of an answer the model reads in the table, at most
# writes the table of the JSON at sys.argv[1], {"rows", "total"}, its first row its heading, as the
# workbook sys.argv[2]: its heading bold, frozen and filtered, its last row bold if a total, its
# amounts numbers, each column as wide as it needs
_SAVE = """
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
book = Workbook()
sheet = book.active
sheet.title = "Answers"
given = json.load(open(sys.argv[1], encoding="utf-8"))
for row in given["rows"]:
    sheet.append(row)
for cell in sheet[1] + (sheet[sheet.max_row] if given["total"] else ()):
    cell.font = Font(bold=True)
for row in sheet.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if isinstance(cell.value, (int, float)):
            cell.number_format = "#,##0.00"
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = sheet.dimensions
for column in sheet.columns:
    width = max(len(str(cell.value or "")) for cell in column)
    sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, width + 2))
book.save(sys.argv[2])
print("{}")
"""
# the system prompt's rule of asking every file
RULE = (
    "For a question of every file, or of many, even of a few, as what the invoices come to or "
    "each one's date, call ask_files once rather than reading them, as it reads every one; answer "
    "with its table, or a long one's first rows, its totals, and the files that gave no answer, "
    "and say what stands out."
)
NAMED = 20  # files a call names at most of those that gave no answer, or could not be read
# a currency, as a symbol or a code, of those Europe's offices write most
_CURRENCY = r"[€£$¥₺₹]|[A-Z]{3}|TL|kr|zł|Kč|Ft|lei|Fr\."
# an amount as a file writes it: a sign, a currency before or after, and digits grouped in
# thousands by . , ' or a space, with cents after a . or a ,
_AMOUNT = re.compile(
    rf"(?P<sign>[-−])?\s*(?P<before>{_CURRENCY})?\s*(?P<minus>[-−])?\s*"
    r"(?P<number>\d{1,3}(?:[.,'\u00a0\u202f ]\d{3})+(?:[.,]\d{1,2})?|\d+(?:[.,]\d{1,2})?)"
    rf"\s*(?P<after>{_CURRENCY})?"
)
# a reader's instructions, of the answers' keys
READING = """\
You read one file for Leat, an assistant, who asks the same of many files. Answer from this file \
alone, as a JSON object of these keys: {keys}. Each value is short: what the file says, a number \
or a date as it is written there, or "" where the file does not say. Write the JSON alone."""


def tools(reader: Client, index: Index | None = None) -> list[Tool]:
    """ask_files, whose readers the model `reader` serves, reading the files' text as `index`
    keeps it, if given."""
    parameters = {
        "type": "object",
        "properties": {
            "question": {"type": "string", "description": "what to find out from each file"},
            "columns": {
                "type": "array", "items": {"type": "string"},
                "description": "the table's columns, each one thing to find, as 'Date' and "
                "'Total'; leave them out for one column, the answer",
            },
            "files": {
                "type": "array", "items": {"type": "string"},
                "description": "the files or folders to ask, by their paths; leave them out to ask "
                "every file",
            },
            "name": {"type": "string", "description": "a short name of what the table holds, as "
                     "'March invoices.xlsx'"},
        },
        "required": ["question", "name"],
    }  # fmt: skip
    return [
        Tool(
            "ask_files",
            "Ask every file in the workspace, or those named, the same question, one by one, as "
            "each invoice's date and total, or which contracts renew themselves. Returns a table "
            "of the answers, a row a file, each numbered, saved as a spreadsheet too",
            parameters,
            lambda context, question, columns=None, files=None, name=None: ask(
                reader, index, context, question, columns, files, name
            ),
            (RULE,),
        )
    ]


def ask(
    reader: Client, index: Index | None, context: Context, question: str,
    columns: list[str] | None = None, files: list[str] | None = None, name: str | None = None,
) -> Result:  # fmt: skip
    space = context.space()
    keys = [" ".join(c.split()) for c in columns or [] if c.strip()] or ["Answer"]
    names = _named_files(space, files)
    if not names:
        return Result("There are no files to ask.", {"question": question, "total": 0})
    if len(names) > FILES:
        raise ValueError(f"a call asks {FILES} files at most, not {len(names)}: name fewer")

    def read(named: str) -> tuple[dict[str, str], str | None] | None:  # none once stopped
        if context.stopped.is_set():
            return None
        return _answer(reader, index, space, named, question, keys, context.person)

    rows: list[tuple[dict[str, str], str | None] | None] = [None] * len(names)
    done = 0
    context.progress({"done": 0, "total": len(names)})
    with concurrent.futures.ThreadPoolExecutor(READERS) as pool:
        futures = {pool.submit(read, named): i for i, named in enumerate(names)}
        for future in concurrent.futures.as_completed(futures):
            if (row := future.result()) is not None:
                rows[futures[future]] = row
                done += 1
                context.progress({"done": done, "total": len(names)})
    asked = [(named, row) for named, row in zip(names, rows, strict=True) if row is not None]
    answered = [(named, answers) for named, (answers, _) in asked]
    unread = [named for named, (_, why) in asked if why]
    unanswered = [named for named, (answers, why) in asked if not why and not any(answers.values())]
    sources = []
    for named, _ in answered:
        n = context.cite(f"file:{named}", named)
        sources.append({"n": n, "url": f"file:{named}", "title": named, "file": named})
    totals = _totals(keys, [row for _, row in answered])
    table: list[list[Any]] = [
        ["File", *(f"{k} ({totals[k]['currency']})" if totals.get(k, {}).get("currency") else k
                   for k in keys)],
        *([named, *(totals[k]["values"][i] if k in totals else row[k] for k in keys)]
          for i, (named, row) in enumerate(answered)),
    ]  # fmt: skip
    if totals:
        table.append(["Total", *(totals[k]["total"] if k in totals else "" for k in keys)])
    saved, why = _save(space, table, bool(totals), name or _named(question))
    shown = [[f"{named} [{s['n']}]" if s["n"] else named, *(row[k][:CELL] for k in keys)]
             for (named, row), s in zip(answered[:SHOWN], sources, strict=False)]  # fmt: skip
    if totals:
        shown.append(["Total", *(totals[k]["said"] if k in totals else "" for k in keys)])
    said = _markdown([["File", *keys], *shown])
    if len(answered) > SHOWN:
        said += f"\n\n(The first {SHOWN} rows of {len(answered)}.)"
    if totals:
        said += f"\n\n(Totals of all {len(answered)} rows, added by code.)"
    if len(answered) < len(names):
        said += f"\n\n(Stopped after {len(answered)} of the {len(names)} files.)"
    if unanswered:
        said += (f"\n\nNo answer in {len(unanswered)} of the {len(answered)} files, which may not "
                 f"say it: {_listed(unanswered)}.")  # fmt: skip
    if unread:
        said += f"\n\nNot read or not asked, as their rows say: {_listed(unread)}."
    said += f"\n\nSaved as {saved}." if saved else f"\n\nIt could not be saved: {why}"
    info = {"question": question, "columns": keys, "done": done, "total": len(names)}
    gaps = {"unanswered": unanswered, "unread": unread}
    return Result(said, info | gaps | {"results": sources} | ({"files": [saved]} if saved else {}))


def _answer(
    reader: Client, index: Index | None, space: Workspace, name: str, question: str,
    keys: list[str], person: int | None,
) -> tuple[dict[str, str], str | None]:  # fmt: skip
    # a file's answers, by their keys, as a reader reads it, an image as it sees it, and None; or
    # if it cannot, why, in the first too
    asked = f"The question: {question}\n\nThe file, {name}"
    content: str | list[dict[str, Any]]
    try:
        if picture(name):
            content = _image(space, name, f"{asked}, is this image.")
        else:
            text = documents.read(space, name, index)
            if not documents.said(text):  # as a scanned PDF's, while no model reads images
                return _row(keys, "(It holds no text: a scan, which only a model that sees images "
                            "reads)")  # fmt: skip
            if len(text) > READ:
                text = _excerpt(index, space, name, text, f"{question} {' '.join(keys)}")
            content = f"{asked}:\n\n{text}"
    except (OSError, ValueError, RuntimeError) as e:
        return _row(keys, f"(It could not be read: {e})")
    system = {
        "role": "system",
        "content": READING.format(keys=json.dumps(keys, ensure_ascii=False)),
    }
    body = {
        "messages": [system, {"role": "user", "content": content}],
        "max_tokens": ANSWER,
        "temperature": 0,
        "reasoning_effort": "none",
    } | BACKGROUND
    try:
        said = reader.reply(body, person)["content"].strip()
    except EngineError as e:
        return _row(keys, f"(It could not be asked: {e})")
    return _parsed(said, keys), None


def _row(keys: list[str], why: str) -> tuple[dict[str, str], str]:
    # a row of a file that could not be asked, why in its first column, and why
    return {k: why if i == 0 else "" for i, k in enumerate(keys)}, why


def _image(space: Workspace, name: str, text: str) -> list[dict[str, Any]]:
    # a message's content of an image of the space's, as a data: URL, and text after it
    kind = mimetypes.guess_type(name)[0] or "image/png"
    url = data_url(space.path(name).read_bytes(), kind)
    return [{"type": "image_url", "image_url": {"url": url}}, {"type": "text", "text": text}]


def _named_files(space: Workspace, named: list[str] | None) -> list[str]:
    # the files a call names, each of its folders' as their names' order has them, or without
    # any named every file of the space's to ask. Raises FileNotFoundError for one that is not
    every = [f["name"] for f in space.files() if _askable(space, f["name"])]
    if not named:
        return sorted(every)
    found: list[str] = []
    for name in named:
        path = space.path(name)
        if path.is_dir():
            folder = path.relative_to(space.root).as_posix() + "/"
            found += sorted(f for f in every if f.startswith(folder) or folder == "./")
        elif path.is_file():
            found.append(path.relative_to(space.root).as_posix())
        else:
            raise space.missing(name)
    return list(dict.fromkeys(found))


def _askable(space: Workspace, name: str) -> bool:
    # whether a file is one to ask: a document, text, or an image, which a reader sees
    return picture(name) or documents.readable(space, name)


def _excerpt(index: Index | None, space: Workspace, name: str, text: str, query: str) -> str:
    # a long file's start, and the passages after it the index finds likeliest to answer the
    # query, in the file's order, READ characters in all at most
    found = index.search(space, query, 20, name) if index else []
    parts, room = [text[:START]], READ - START
    for passage in sorted(found, key=lambda p: p["start"]):
        if passage["start"] >= START and len(passage["text"]) < room:
            parts.append(passage["text"])
            room -= len(passage["text"])
    return "\n\n[…]\n\n".join(parts)


def _parsed(said: str, keys: list[str]) -> dict[str, str]:
    # a reader's answers, by their keys, of the JSON object it wrote, found by its keys whatever
    # their case; of an answer that is not one, the whole of it in the first
    try:
        found = json.loads(said[said.index("{") : said.rindex("}") + 1])
        assert isinstance(found, dict)
    except (ValueError, AssertionError):
        return {k: " ".join(said.split()) if i == 0 else "" for i, k in enumerate(keys)}
    folded = {str(k).casefold(): v for k, v in found.items()}
    answers = {}
    for key in keys:
        value = folded.get(key.casefold(), "")
        answers[key] = " ".join(str(value if value is not None else "").split())
    return answers


def _save(
    space: Workspace, table: list[list[Any]], total: bool, name: str
) -> tuple[str | None, str]:
    # saves a table as a spreadsheet of the space's, its last row bold if a `total`, by a name
    # free there, in the sandbox, with its libraries; returns its name, or None and why not
    if space.environment is None:
        return None, "the sandbox has no spreadsheet library"
    name = space.new(name if name.lower().endswith(".xlsx") else f"{name}.xlsx")
    try:
        space.given(_SAVE, {"rows": table, "total": total}, name)
    except ValueError as e:
        return None, str(e)
    return name, ""


def _named(question: str) -> str:
    # a spreadsheet's name of a question's first words
    words = re.sub(r"[^\w\s'-]", "", question).split()[:6]
    return " ".join(words).capitalize() or "Answers"


def _markdown(rows: list[list[str]]) -> str:
    # rows as a markdown table, the first its heading
    def cell(text: str) -> str:
        return " ".join(text.split()).replace("|", "\\|") or " "

    lines = ["| " + " | ".join(cell(c) for c in row) + " |" for row in rows]
    return "\n".join([lines[0], "|" + " --- |" * len(rows[0]), *lines[1:]])


def _totals(keys: list[str], rows: list[dict[str, str]]) -> dict[str, dict[str, Any]]:
    # the totals of the columns of amounts, by their keys: a column of two values at least, each
    # an amount, of one currency, and half of them with a currency or cents, as an amount and not
    # an invoice's number or a count is written. Each total's value, as it is said, in the
    # column's way of writing amounts, its currency, and each row's value, or None if blank
    totals = {}
    for key in keys:
        said = [row[key] for row in rows]
        amounts = [_amount(text) for text in said if text]
        if len(amounts) < 2 or None in amounts:
            continue
        found = [a for a in amounts if a is not None]  # none, but as mypy knows it
        currencies = {a["currency"] for a in found if a["currency"]}
        if len(currencies) > 1 or 2 * sum(a["marked"] for a in found) < len(found):
            continue
        commas = sum(a["decimal"] == "," for a in found) > sum(a["decimal"] == "." for a in found)
        total = round(sum(a["value"] for a in found), 2)
        first = next((a for a in found if a["currency"]), found[0])
        number = f"{abs(total):,.2f}"
        if commas:
            number = number.replace(",", " ").replace(".", ",").replace(" ", ".")
        currency, before = first["currency"], first["before"]
        written = f"{currency}{number}" if before and not currency.isalpha() else (
            f"{currency} {number}" if before else f"{number} {currency}".strip())  # fmt: skip
        values = iter(a["value"] for a in found)
        totals[key] = {
            "total": total, "said": f"-{written}" if total < 0 else written,
            "currency": currency, "values": [next(values) if text else None for text in said],
        }  # fmt: skip
    return totals


def _amount(text: str) -> dict[str, Any] | None:
    # an amount, as _AMOUNT reads it: its value, its currency, whether the currency was before it,
    # its decimal separator, if it has cents, and whether a currency or cents mark it as one
    match = _AMOUNT.fullmatch(text.strip())
    if not match:
        return None
    number = match["number"]
    cents = re.search(r"[.,](\d{1,2})$", number)
    whole = int(re.sub(r"\D", "", number[: cents.start()] if cents else number))
    value = whole + (int(cents[1]) / 10 ** len(cents[1]) if cents else 0)
    currency = match["before"] or match["after"] or ""
    return {
        "value": -value if match["sign"] or match["minus"] else value, "currency": currency,
        "before": bool(match["before"]), "decimal": number[cents.start()] if cents else "",
        "marked": bool(currency) or bool(cents and len(cents[1]) == 2),
    }  # fmt: skip


def _listed(names: list[str]) -> str:
    # files by their names, NAMED at most, and how many more
    more = f", and {len(names) - NAMED} more" if len(names) > NAMED else ""
    return ", ".join(names[:NAMED]) + more

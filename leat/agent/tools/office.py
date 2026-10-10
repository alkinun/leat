"""Documents made and changed as an office makes and changes them, as tools: make_document makes a
Word document or a PDF of the model's markdown, fill_template fills a template's fields, and
suggest_edits suggests edits as tracked changes, which the user accepts or rejects in Word. Each
works in the sandbox, with python-docx or fpdf2, so that a document made to attack its parser
attacks nothing else; a template or a document changed is read there, and left as it was.

A document made is the model's writing alone, as markdown: its headings, paragraphs, lists, quotes,
code and tables, and **bold**, *italic* and `code` within them, each line of a paragraph a line,
as a letter's address is written. The harness lays it out, so that a model that writes well makes
a document without writing code. A Word document takes a template's styles, headers and footers,
as the office's letterhead, its body left out; or is in Leat's own style, as a PDF is, in DejaVu,
which has every letter.

Filled and edited documents keep their formatting: new text takes the formatting of the first
character of what it replaces, however Word split that among its runs. A field is written
"{{Client name}}", or as Word shows a merge field, "«Client name»"; values are matched to fields
whatever their case and spaces. An edit is of a passage written once in the document, within a
paragraph, which it deletes and puts its replacement after, both marked as Leat's, with a comment
if it has one.
"""

import re
from pathlib import PurePosixPath
from typing import Any

from leat.agent.tools import Context, Result, Tool
from leat.agent.workspace import Workspace

# the paragraphs of the document at sys.argv[2], its body's, its tables', its headers' and
# footers', in which a function `change` of each paragraph's text and of the spans of its runs
# changes the text; defined for the scripts below, which define `change`
_PARAGRAPHS = """
import json, re, sys
from docx import Document

def paragraphs(document):
    def within(parent):
        for paragraph in parent.paragraphs:
            yield paragraph
        for table in getattr(parent, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from within(cell)
    yield from within(document)
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header,
                     section.first_page_footer, section.even_page_header, section.even_page_footer):
            if not part.is_linked_to_previous:
                yield from within(part)

def replace(paragraph, start, end, new):
    # replaces the paragraph's text from start to end, over its runs, with new, in the run where
    # it starts, which keeps that run's formatting
    runs, at = paragraph.runs, 0
    spans = []
    for run in runs:
        spans.append((at, at + len(run.text)))
        at += len(run.text)
    first = next(i for i, (a, b) in enumerate(spans) if a <= start < b)
    last = next(i for i, (a, b) in enumerate(spans) if a < end <= b)
    head = runs[first].text[: start - spans[first][0]]
    tail = runs[last].text[end - spans[last][0]:]
    for i in range(first + 1, last):
        runs[i].text = ""
    if first == last:
        runs[first].text = head + new + tail
    else:
        runs[first].text, runs[last].text = head + new, tail
"""
# fills the fields of the document at sys.argv[2] with the values of the JSON at sys.argv[1], and
# saves it at sys.argv[3]; prints the fields filled, those without a value, and the values of no
# field, as JSON
_FILL = (
    _PARAGRAPHS
    + """
given = json.load(open(sys.argv[1]))
values = {" ".join(k.split()).casefold(): str(v) for k, v in given.items()}
document = Document(sys.argv[2])
FIELD = re.compile(r"\\{\\{\\s*([^{}]+?)\\s*\\}\\}|«\\s*([^«»]+?)\\s*»")
filled, missing = set(), set()
for paragraph in paragraphs(document):
    text = "".join(run.text for run in paragraph.runs)
    for match in reversed(list(FIELD.finditer(text))):
        name = " ".join((match[1] or match[2]).split())
        if name.casefold() in values:
            replace(paragraph, match.start(), match.end(), values[name.casefold()])
            filled.add(name)
        else:
            missing.add(name)
document.save(sys.argv[3])
used = {name.casefold() for name in filled}
unused = sorted(k for k in given if " ".join(k.split()).casefold() not in used)
print(json.dumps({"filled": sorted(filled), "missing": sorted(missing - filled), "unused": unused},
                 ensure_ascii=False))
"""
)


# suggests the edits of the JSON at sys.argv[1], each {"find", "replace", "comment"}, to the
# document at sys.argv[2] as tracked changes, and saves it at sys.argv[3]; prints how many times
# each edit's passage was found, as JSON, those found once made
_SUGGEST = (
    _PARAGRAPHS
    + """
import copy, datetime
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run

edits = json.load(open(sys.argv[1]))
document = Document(sys.argv[2])
when = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
taken = [int(v) for e in document.element.iter() if (v := e.get(qn("w:id"))) and v.isdigit()]
ids = iter(range(max(taken, default=0) + 1, 1 << 31))

def split(paragraph, at):
    # splits the run that holds the paragraph's character `at` there, so that a run begins at it
    position = 0
    for run in paragraph.runs:
        n = len(run.text)
        if position < at < position + n:
            text, twin = run.text, copy.deepcopy(run._r)
            run._r.addnext(twin)
            run.text, Run(twin, paragraph).text = text[: at - position], text[at - position:]
            return
        position += n

def covering(paragraph, start, end):
    # the runs that hold the paragraph's text from start to end, and no more
    split(paragraph, start)
    split(paragraph, end)
    runs, position = [], 0
    for run in paragraph.runs:
        if start <= position < end and run.text:
            runs.append(run)
        position += len(run.text)
    return runs

def change(tag):
    # a tracked change of Leat's, made now
    element = OxmlElement(tag)
    for key, value in (("w:id", str(next(ids))), ("w:author", "Leat"), ("w:date", when)):
        element.set(qn(key), value)
    return element

found = []
for edit in edits:
    find, new, note = edit["find"], edit.get("replace") or "", edit.get("comment")
    places = []
    for paragraph in paragraphs(document):
        text, at = "".join(run.text for run in paragraph.runs), 0
        while find and (at := text.find(find, at)) >= 0:
            places.append((paragraph, at))
            at += len(find)
    found.append(len(places))
    if len(places) != 1:
        continue
    paragraph, start = places[0]
    runs = covering(paragraph, start, start + len(find))
    deleted = change("w:del")
    runs[0]._r.addprevious(deleted)
    for run in runs:
        for t in run._r.findall(qn("w:t")):
            t.tag = qn("w:delText")
        deleted.append(run._r)
    marked = runs
    if new:
        inserted, r = change("w:ins"), copy.deepcopy(runs[0]._r)
        for child in [c for c in r if c.tag != qn("w:rPr")]:
            r.remove(child)
        inserted.append(r)
        deleted.addnext(inserted)
        Run(r, paragraph).text = new
        marked = [Run(r, paragraph)]
    if note:
        document.add_comment(marked, note, author="Leat", initials="L")
document.save(sys.argv[3])
print(json.dumps(found))
"""
)

# makes a Word document of the blocks of the JSON at sys.argv[1], as blocks() reads them, at
# sys.argv[2], in the styles, headers and footers of the template at sys.argv[3], if one, its body
# left out, or in Leat's: Calibri of 11 points, in margins of 2 cm. A style the template lacks, as
# "Heading 2", is stood for by formatting
_WORD = """
import json, sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

blocks = json.load(open(sys.argv[1], encoding="utf-8"))
if len(sys.argv) > 3:
    document = Document(sys.argv[3])
    body = document.element.body
    for child in list(body):
        if not child.tag.endswith("}sectPr"):
            body.remove(child)
else:
    document = Document()
    document.styles["Normal"].font.name, document.styles["Normal"].font.size = "Calibri", Pt(11)
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin = section.bottom_margin = Cm(2)
styles = {style.name for style in document.styles}

def paragraph(spans, style=None, bold=False, italic=False, size=None, font=None):
    p = document.add_paragraph(style=style if style in styles else None)
    for text, marks in spans:
        for i, line in enumerate(text.split("\\n")):
            if i:
                run.add_break()  # the line before's
            run = p.add_run(line)
            run.bold = True if bold or "b" in marks else None
            run.italic = True if italic or "i" in marks else None
            if size:
                run.font.size = Pt(size)
            if font or "c" in marks:
                run.font.name = font or "Consolas"
    return p

for block in blocks:
    kind = block[0]
    if kind == "heading":
        style = f"Heading {block[1]}"
        size = {1: 16, 2: 13}.get(block[1], 11) if style not in styles else None
        paragraph(block[2], style, bold=style not in styles, size=size)
    elif kind == "paragraph":
        paragraph(block[1])
    elif kind == "item":
        depth, label, spans = block[1:]
        p = paragraph([[f"{label}\\t", ""], *spans])
        indent = Cm(0.63 * (depth + 1))
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = indent, -Cm(0.63)
        p.paragraph_format.tab_stops.add_tab_stop(indent)
        p.paragraph_format.space_after = Pt(2)
    elif kind == "quote":
        p = paragraph(block[1], "Quote", italic="Quote" not in styles)
        if "Quote" not in styles:
            p.paragraph_format.left_indent = Cm(1)
    elif kind == "code":
        paragraph([[block[1], ""]], font="Consolas", size=9.5)
    elif kind == "table":
        head, rows, aligns = block[1:]
        table = document.add_table(rows=1 + len(rows), cols=len(head))
        if "Table Grid" in styles:
            table.style = "Table Grid"
        for i, row in enumerate([head, *rows]):
            for j in range(len(head)):
                cell = table.cell(i, j).paragraphs[0]
                cell.alignment = getattr(WD_ALIGN_PARAGRAPH, aligns[j])
                for text, marks in row[j] if j < len(row) else []:
                    run = cell.add_run(text)
                    run.bold = True if i == 0 or "b" in marks else None
                    run.italic = True if "i" in marks else None
        document.add_paragraph()
    elif kind == "rule":
        document.add_paragraph()
document.save(sys.argv[2])
print("{}")
"""
# makes a PDF of the blocks of the JSON at sys.argv[1] at sys.argv[2], on A4, in margins of 2 cm, in
# DejaVu, which matplotlib has
_PDF = """
import json, os, sys
import matplotlib
from fpdf import FPDF

blocks = json.load(open(sys.argv[1], encoding="utf-8"))
fonts = os.path.join(matplotlib.get_data_path(), "fonts", "ttf")
pdf = FPDF(format="A4")
pdf.set_margins(20, 20, 20)
pdf.set_auto_page_break(True, 20)
for style, name in (("", ""), ("B", "-Bold"), ("I", "-Oblique"), ("BI", "-BoldOblique")):
    pdf.add_font("Sans", style, os.path.join(fonts, f"DejaVuSans{name}.ttf"))
pdf.add_font("Mono", "", os.path.join(fonts, "DejaVuSansMono.ttf"))
pdf.add_page()
SIZE = 10.5

def write(spans, size=SIZE, bold=False, italic=False):
    for text, marks in spans:
        if "c" in marks:
            pdf.set_font("Mono", "", size * 0.9)
        else:
            b, i = bold or "b" in marks, italic or "i" in marks
            pdf.set_font("Sans", ("B" if b else "") + ("I" if i else ""), size)
        pdf.write(size * 0.5, text)
    pdf.ln(size * 0.5)

for block in blocks:
    kind = block[0]
    if kind == "heading":
        pdf.ln(3)
        write(block[2], {1: 16, 2: 13}.get(block[1], 11), bold=True)
        pdf.ln(1)
    elif kind == "paragraph":
        write(block[1])
        pdf.ln(2.5)
    elif kind == "item":
        depth, label, spans = block[1:]
        indent = 20 + 6 * (depth + 1)
        pdf.set_x(indent - 5)
        pdf.set_font("Sans", "", SIZE)
        pdf.cell(5, SIZE * 0.5, label)
        pdf.set_left_margin(indent)
        write(spans)
        pdf.set_left_margin(20)
        pdf.ln(1)
    elif kind == "quote":
        pdf.ln(1.5)
        pdf.set_left_margin(28)
        pdf.set_x(28)
        pdf.set_text_color(80)
        write(block[1], italic=True)
        pdf.set_text_color(0)
        pdf.set_left_margin(20)
        pdf.ln(2.5)
    elif kind == "code":
        pdf.set_font("Mono", "", SIZE * 0.85)
        pdf.multi_cell(0, SIZE * 0.45, block[1])
        pdf.ln(2.5)
    elif kind == "table":
        head, rows, aligns = block[1:]
        pdf.ln(1)
        pdf.set_font("Sans", "", SIZE)
        aligned = tuple(aligns)
        with pdf.table(text_align=aligned, line_height=SIZE * 0.55, padding=(0.8, 1.5)) as table:
            for row in [head, *rows]:
                cells = table.row()
                for j in range(len(head)):
                    cells.cell("".join(text for text, _ in row[j]) if j < len(row) else "")
        pdf.ln(2.5)
    elif kind == "rule":
        y = pdf.get_y() + 1
        pdf.line(20, y, 190, y)
        pdf.ln(4)
pdf.output(sys.argv[2])
print("{}")
"""


def tools() -> list[Tool]:
    """make_document, fill_template and suggest_edits, in their call's conversation's
    workspace."""
    return [
        Tool(
            "make_document",
            "Make a Word document or a PDF, as a letter, memo or report, of its text in markdown: "
            "headings, paragraphs, lists, quotes and tables, and **bold** and *italic*, each line "
            "of a paragraph a line, as an address. Laid out for you, in the office's letterhead "
            "if a Word template is given. Makes a new file",
            {
                "type": "object",
                "properties": {
                    "name": {
                        "type": "string",
                        "description": "the document's name, as 'Letter - Hartley.docx', or "
                        "'Letter - Hartley.pdf' for a PDF",
                    },
                    "content": {"type": "string", "description": "the document, in markdown"},
                    "template": {
                        "type": "string",
                        "description": "a Word document whose styles, headers and footers the "
                        "document takes, as the office's letterhead, its text left out; leave it "
                        "out for Leat's own style",
                    },
                },
                "required": ["name", "content"],
            },  # fmt: skip
            lambda context, name, content, template=None: make(context, name, content, template),
        ),
        Tool(
            "fill_template",
            "Fill a Word template's fields, written {{Client name}} or «Client name», keeping its "
            "formatting. Makes a new document, the template unchanged, and says which fields had "
            "no value. Call it once for each document",
            {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "the template, a .docx file"},
                    "values": {
                        "type": "object",
                        "additionalProperties": {"type": "string"},
                        "description": "each field's value, by its name",
                    },
                    "name": {
                        "type": "string",
                        "description": "the new document's name, as 'Letter - Hartley.docx'",
                    },
                },
                "required": ["path", "values", "name"],
            },  # fmt: skip
            lambda context, path, values, name: fill(context, path, values, name),
        ),
        Tool(
            "suggest_edits",
            "Suggest edits to a Word document as tracked changes, which the user accepts or "
            "rejects in Word, each with a comment that says why. Makes a new document, the "
            "original unchanged. Each passage must be written exactly once in the document, "
            "within one paragraph",
            {
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "the document, a .docx file"},
                    "edits": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "find": {
                                    "type": "string",
                                    "description": "the passage, exactly as it is written there",
                                },
                                "replace": {
                                    "type": "string",
                                    "description": "the text that replaces it; empty to delete it",
                                },
                                "comment": {"type": "string", "description": "why"},
                            },
                            "required": ["find", "replace"],
                        },
                    },
                    "name": {
                        "type": "string",
                        "description": "the new document's name, as 'Lease - suggested.docx'",
                    },
                },
                "required": ["path", "edits", "name"],
            },  # fmt: skip
            lambda context, path, edits, name: suggest(context, path, edits, name),
        ),
    ]


def make(context: Context, name: str, content: str, template: str | None = None) -> Result:
    space = context.space()
    pdf = PurePosixPath(name).suffix.lower() == ".pdf"
    if pdf and template:
        raise ValueError("a template's letterhead is kept in a Word document alone: name it .docx")
    laid = blocks(content)
    if not laid:
        raise ValueError("the document has no text: give its content, in markdown")
    letterhead = _docx(space, template) if template else None
    made = space.free(name if pdf else _named(name), folders=True)
    space.path(made).parent.mkdir(parents=True, exist_ok=True)
    space.given(_PDF if pdf else _WORD, laid, made, *([letterhead] if letterhead else []))
    styled = f", in the styles, headers and footers of {letterhead}" if letterhead else ""
    return Result(f"Made {made}{styled}.", {"files": [made]})


def fill(context: Context, path: str, values: dict[str, Any], name: str) -> Result:
    space = context.space()
    template = _docx(space, path)
    made = space.free(_named(name), folders=True)
    space.path(made).parent.mkdir(parents=True, exist_ok=True)
    said = space.given(_FILL, values, template, made)
    filled, missing, unused = said["filled"], said["missing"], said["unused"]
    lines = [f"Made {made}, of {template}: {len(filled)} fields filled."]
    if missing:
        lines.append(f"Left as they were, without values: {', '.join(missing)}.")
    if unused:
        lines.append(f"Values of no field there: {', '.join(unused)}.")
    if not filled and not missing:
        lines.append("The template has no fields, written {{Name}} or «Name».")
    return Result(" ".join(lines), {"files": [made]} | said)


def suggest(context: Context, path: str, edits: list[dict[str, Any]], name: str) -> Result:
    space = context.space()
    document = _docx(space, path)
    edits = [_edit(e) for e in edits] if isinstance(edits, list) else []
    if not edits or not all(e["find"] for e in edits):
        raise ValueError('edits must be a list of objects, each {"find": "the passage, as it is '
                         'written", "replace": "what replaces it", "comment": "why"}')  # fmt: skip
    made = space.free(_named(name), folders=True)
    space.path(made).parent.mkdir(parents=True, exist_ok=True)
    found = space.given(_SUGGEST, edits, document, made)
    done = sum(n == 1 for n in found)
    lines = [f"Made {made}, of {document}: {done} of {len(edits)} edits suggested, as tracked "
             "changes, which the user accepts or rejects in Word."]  # fmt: skip
    for edit, n in zip(edits, found, strict=True):
        if n != 1:
            where = "is not in it" if n == 0 else f"is in it {n} times: give more of it"
            lines.append(f"Not made: “{edit['find'][:80]}” {where}.")
    return Result("\n".join(lines), {"files": [made], "suggested": done, "edits": len(edits)})


def _edit(given: Any) -> dict[str, str]:
    # an edit as suggest_edits takes it, of the names a model may give its parts: the passage, as
    # "find" or "old", its replacement, as "replace" or "new", and why, as "comment" or "reason"
    if not isinstance(given, dict):
        return {"find": ""}
    first = lambda *keys: next((str(given[k]) for k in keys if given.get(k)), "")  # noqa: E731
    return {"find": first("find", "old", "passage", "text", "original"),
            "replace": first("replace", "new", "replacement", "with"),
            "comment": first("comment", "reason", "why", "note")}  # fmt: skip


def _docx(space: Workspace, path: str) -> str:
    # a Word document's name in the space. Raises FileNotFoundError, or ValueError of another kind
    file = space.path(path)
    if not file.is_file():
        raise space.missing(path)
    if file.suffix.lower() != ".docx":
        raise ValueError(f"{path} is not a Word document, a .docx")
    return file.relative_to(space.root).as_posix()


def _named(name: str) -> str:
    # a name of a Word document, as given, its kind added, or in place of another document's
    path = PurePosixPath(name)
    if path.suffix.lower() in (".pdf", ".doc", ".odt", ".rtf", ".txt"):
        return str(path.with_suffix(".docx"))
    return name if path.suffix.lower() == ".docx" else f"{name}.docx"


def blocks(text: str) -> list[list[Any]]:
    """The blocks of markdown, as make_document lays them out, each a list of its kind and parts:
    ["heading", level, spans], ["paragraph", spans], ["item", depth, label, spans], ["quote",
    spans], ["code", text], ["table", heading's cells, rows of cells, each cell spans, each column's
    alignment] or ["rule"]. A paragraph's lines are lines; a list's items are labelled as written,
    3. or •; a column is aligned LEFT, RIGHT or CENTER as its divider says, :--, --: or :-:; and a
    document a model fenced whole as markdown is read within its fence."""
    text = re.sub(r"<br\s*/?>", "\n", text.replace("\r\n", "\n"))
    if fenced := _FENCED.fullmatch(text):
        text = fenced[2]
    lines, i = text.split("\n"), 0
    out: list[list[Any]] = []
    levels: list[int] = []  # the indents of the list items before, by their depths
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if not _ITEM.match(line):
            levels = []
        if fence := _FENCE.match(line):
            body, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith(fence[1]):
                body.append(lines[i])
                i += 1
            out.append(["code", "\n".join(body)])
            i += 1
        elif heading := _HEADING.match(line):
            out.append(["heading", len(heading[1]), spans(heading[2])])
            i += 1
        elif _RULE.match(line):
            out.append(["rule"])
            i += 1
        elif "|" in line and i + 1 < len(lines) and _DIVIDER.match(lines[i + 1]):
            head, rows, i = _cells(line), [], i + 2
            divider = i - 1
            while i < len(lines) and "|" in lines[i]:
                rows.append([spans(cell) for cell in _cells(lines[i])])
                i += 1
            columns = _cells(lines[divider]) + [""] * len(head)
            aligns = [_ALIGNS[c.startswith(":"), c.endswith(":")] for c in columns[: len(head)]]
            out.append(["table", [spans(cell) for cell in head], rows, aligns])
        elif item := _ITEM.match(line):
            indent = len(item[1].expandtabs(4))
            while levels and indent < levels[-1]:
                levels.pop()
            if not levels or indent > levels[-1]:
                levels.append(indent)
            depth, said, i = len(levels) - 1, [item[3]], i + 1
            while i < len(lines) and lines[i].strip() and not _starts(lines, i):  # its lines after
                said.append(lines[i].strip())
                i += 1
            label = item[2] if item[2][0].isdigit() else "•–◦"[min(depth, 2)]
            out.append(["item", depth, label, spans("\n".join(said))])
        elif line.lstrip().startswith(">"):
            said = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                said.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            out.append(["quote", spans("\n".join(said).strip())])
        else:
            said = []
            while i < len(lines) and lines[i].strip() and (not said or not _starts(lines, i)):
                said.append(lines[i].strip())
                i += 1
            out.append(["paragraph", spans("\n".join(said))])
    return out


def spans(text: str) -> list[list[str]]:
    """Markdown's inline text as spans, each its text and marks: "b" bold, "i" italic and "c"
    code. A ** or * opens before a non-space and closes after one, an _ neither within a word, and
    one that nothing closes is text; a link is its text, and its address if the text is not."""
    tokens: list[list[Any]] = []  # each [text] or [delimiter, can open, can close]
    for match in _INLINE.finditer(text):
        escaped, code, link, address, delimiter, plain = match.groups()
        if delimiter:
            before = text[match.start() - 1 : match.start()]
            after = text[match.end() : match.end() + 1]
            opens, closes = bool(after.strip()), bool(before.strip())
            if delimiter[0] == "_" and before.isalnum() and after.isalnum():
                opens = closes = False
            tokens.append([delimiter, opens, closes])
        elif code is not None:
            tokens.append([code, "c"])
        elif link is not None:
            tokens.append([link if address in link else f"{link} ({address})"])
        else:
            tokens.append([escaped or plain])
    open_: dict[str, int] = {}  # of each delimiter, the token that opened it, unclosed
    paired: set[int] = set()
    for n, token in enumerate(tokens):
        if len(token) == 3:
            kind = token[0].replace("_", "*")
            if token[2] and kind in open_:
                paired.update((open_.pop(kind), n))
            elif token[1]:
                open_[kind] = n
    out: list[list[str]] = []
    marks = ""
    for n, token in enumerate(tokens):
        if len(token) == 3 and n in paired:
            mark = "b" if len(token[0]) == 2 else "i"
            marks = marks.replace(mark, "") if mark in marks else marks + mark
            continue
        said, mark = token[0], token[1] if len(token) == 2 else ""
        if out and out[-1][1] == marks + mark:
            out[-1][0] += said
        else:
            out.append([said, marks + mark])
    return [span for span in out if span[0]]


_FENCE = re.compile(r" {0,3}(`{3,}|~{3,})")
_FENCED = re.compile(r"\s*(`{3,}|~{3,})\s*(?:markdown|md)?\s*\n(.*?)\n\s*\1\s*", re.S)  # whole
_HEADING = re.compile(r" {0,3}(#{1,6})[ \t]+(.*?)(?:[ \t]+#+)?[ \t]*$")
_RULE = re.compile(r" {0,3}([-*_])(?:[ \t]*\1){2,}[ \t]*$")
_ITEM = re.compile(r"(\s*)([-*+•]|\d{1,9}[.)])[ \t]+(.*)")
# a column's alignment, of whether its divider starts and ends with a colon
_ALIGNS = {
    (False, False): "LEFT", (True, False): "LEFT", (False, True): "RIGHT", (True, True): "CENTER",
}  # fmt: skip
_DIVIDER = re.compile(r"\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)*\|?\s*$")
# an escape, a code span, a link, a run of ** or *, or of __ or _, or text
_INLINE = re.compile(
    r"\\([!-/:-@\[-`{-~])|`([^`]+)`|\[([^\]]+)\]\(([^)\s]+)\)|(\*\*|\*|__|_)|([^\\`\[*_]+|.)"
)


def _starts(lines: list[str], i: int) -> bool:
    # whether line i starts a block other than a paragraph
    line = lines[i]
    return bool(
        _FENCE.match(line) or _HEADING.match(line) or _RULE.match(line) or _ITEM.match(line)
        or line.lstrip().startswith(">")
        or ("|" in line and i + 1 < len(lines) and _DIVIDER.match(lines[i + 1]))
    )  # fmt: skip


def _cells(row: str) -> list[str]:
    # a table's row's cells, without its outer pipes; \| a pipe within one
    row = row.strip().removeprefix("|")
    row = row[:-1] if row.endswith("|") and not row.endswith("\\|") else row
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", row)]

"""The workspace's files, as tools: reading them, writing and editing them, and running Python among
them in the sandbox. read reads the skills too, the guides to making documents, at skills/."""

from pathlib import Path
from typing import Any

from leat.agent.tools import Result, Tool, strings
from leat.agent.workspace import Workspace

READ = 12000  # characters of a file read at once
SKILLS = Path(__file__).parent.parent / "skills"
LIBRARIES = "python-docx, openpyxl, python-pptx, fpdf2, pypdf, matplotlib, pandas"
# the kinds of file read by parsing them, in the sandbox, and those that cannot be read as text
DOCUMENTS = (".pdf", ".docx", ".xlsx", ".pptx")
IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".heic")
# reads the document at sys.argv[1] as markdown, which keeps its headings, lists and tables, as
# markitdown reads one, in the sandbox, with its libraries
_EXTRACT = """
import sys
path = sys.argv[1]
kind = path.rsplit(".", 1)[-1].lower()

def table(rows):
    rows = [["" if cell is None else " ".join(str(cell).split()) for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    print()
    for i, row in enumerate(rows):
        print("| " + " | ".join(row + [""] * (width - len(row))) + " |")
        if i == 0:
            print("|" + " --- |" * width)
    print()

if kind == "pdf":
    from pypdf import PdfReader
    for i, page in enumerate(PdfReader(path).pages, 1):
        print(f"## Page {i}\\n\\n{page.extract_text() or ''}\\n")
elif kind == "docx":
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(path)
    for block in document.element.body.iterchildren():
        if block.tag.endswith("}tbl"):
            table([[cell.text for cell in row.cells] for row in Table(block, document).rows])
        elif block.tag.endswith("}p"):
            paragraph = Paragraph(block, document)
            style, text = paragraph.style.name if paragraph.style else "", paragraph.text
            if not text.strip():
                continue
            if style == "Title":
                print(f"# {text}\\n")
            elif style.startswith("Heading") and style[-1:].isdigit():
                print("#" * (int(style[-1]) + 1) + f" {text}\\n")
            elif style.startswith("List Number"):
                print(f"1. {text}")
            elif style.startswith("List"):
                print(f"- {text}")
            else:
                print(f"{text}\\n")
elif kind == "xlsx":  # each cell's value, or its formula if no program has computed it yet
    from openpyxl import load_workbook
    values, formulas = (load_workbook(path, read_only=True, data_only=d) for d in (True, False))
    for sheet, written in zip(values, formulas):
        print(f"## {sheet.title}\\n")
        table([[raw[i] if value is None else value for i, value in enumerate(row)]
               for row, raw in zip(sheet.iter_rows(values_only=True),
                                   written.iter_rows(values_only=True))])
elif kind == "pptx":
    from pptx import Presentation
    for i, slide in enumerate(Presentation(path).slides, 1):
        title = slide.shapes.title.text if slide.shapes.title is not None else ""
        print(f"## Slide {i}: {title}\\n" if title else f"## Slide {i}\\n")
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_table:
                table([[cell.text for cell in row.cells] for row in shape.table.rows])
            elif shape.has_text_frame and shape.text_frame.text.strip():
                print(shape.text_frame.text + "\\n")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            print(f"Notes: {slide.notes_slide.notes_text_frame.text}\\n")
"""


def tools(workspace: Workspace) -> list[Tool]:
    start = {"type": "integer", "description": "the character to read from, in a long file"}
    path = {"type": "string", "description": "the file's path in the workspace, as 'notes.txt'"}
    reading = {"type": "object", "properties": {"path": path, "start": start}, "required": ["path"]}
    return [
        Tool(
            "read",
            "Read a file in the workspace as text, a PDF or a Word, Excel or PowerPoint file too, "
            "or list a folder's files, '.' the workspace's",
            reading,
            lambda context, path, start=0: read(workspace, path, start),
        ),
        Tool(
            "write",
            "Write a text file in the workspace, in place of any of its name",
            strings(path="the file's path in the workspace", content="the text"),
            lambda context, path, content: write(workspace, path, content),
        ),
        Tool(
            "edit",
            "Replace a passage of a text file in the workspace",
            strings(
                path="the file's path in the workspace",
                old="the passage, as it is in the file",
                new="what replaces it",
            ),  # fmt: skip
            lambda context, path, old, new: edit(workspace, path, old, new),
        ),
        Tool(
            "run",
            "Run Python code in the workspace, in a sandbox without the network: what it prints, "
            "and the files it makes or changes",
            strings(code=f"Python 3; its libraries are {LIBRARIES}"),
            lambda context, code: run(workspace, code),
        ),
    ]


def read(workspace: Workspace, path: str, start: int = 0) -> Result:
    file = _skill(path) if path.startswith("skills/") else workspace.path(path)
    if file.is_dir():
        names = sorted(f"{p.name}/" if p.is_dir() else p.name for p in file.iterdir())
        listed = "\n".join(name for name in names if not name.startswith("."))
        return Result(listed or "The folder is empty.", {"file": path})
    if not file.is_file():
        raise FileNotFoundError(f"there is no file {path}")
    suffix = file.suffix.lower()
    if suffix in IMAGES:
        raise ValueError(f"{path} is an image, which you cannot see")
    if suffix in DOCUMENTS:
        name = f"/workspace/{file.relative_to(workspace.root).as_posix()}"
        ran = workspace.run(_EXTRACT, name, kept=None)  # read on in parts, below
        if ran.status != 0:
            why = (ran.output.strip().splitlines() or ["it stopped"])[-1]
            raise ValueError(f"{path} could not be read: {why}")
        text = ran.output
    else:
        data = file.read_bytes()
        if b"\0" in data[:8192]:
            raise ValueError(f"{path} is not a file of text")
        text = data.decode("utf-8", "replace")
    part = text[start : start + READ]
    if (end := start + len(part)) < len(text):
        part += f"\n\n(characters {start} to {end} of {len(text)}; read on from start={end})"
    return Result(part or "The file is empty.", {"file": path})


def write(workspace: Workspace, path: str, content: str) -> Result:
    file = workspace.path(path)
    file.parent.mkdir(parents=True, exist_ok=True)
    file.write_text(content)
    return Result(f"Wrote {path}, {len(content)} characters.", {"files": [_name(workspace, file)]})


def edit(workspace: Workspace, path: str, old: str, new: str) -> Result:
    file = workspace.path(path)
    if not file.is_file():
        raise FileNotFoundError(f"there is no file {path}")
    text = file.read_text()
    if (n := text.count(old)) != 1:
        found = "is not in" if n == 0 else f"is in {n} places of"
        raise ValueError(f"the passage {found} {path}: give it as it is, enough of it to be one")
    file.write_text(text.replace(old, new))
    return Result(f"Edited {path}.", {"files": [_name(workspace, file)]})


def run(workspace: Workspace, code: str) -> Result:
    before = _state(workspace)
    ran = workspace.run(code)
    after = _state(workspace)
    changed = [name for name, state in after.items() if before.get(name) != state]
    lines = [ran.output.rstrip() or "It printed nothing."]
    if ran.status is None:
        lines.append("It ran out of time, and was stopped.")
    elif ran.status != 0:
        lines.append(f"It failed, with exit status {ran.status}.")
    if changed:
        lines.append(f"Files made or changed: {', '.join(changed)}")
    return Result("\n".join(lines), {"files": changed, "status": ran.status})


def skills() -> list[tuple[str, str]]:
    """Each skill's path, to read, and what it is for, from its SKILL.md's front matter."""
    found = []
    for file in sorted(SKILLS.glob("*/SKILL.md")):
        lines = file.read_text().split("---")[1].strip().splitlines()
        about = dict(line.split(":", 1) for line in lines if ":" in line)
        found.append((f"skills/{file.parent.name}/SKILL.md", about["description"].strip()))
    return found


def _skill(path: str) -> Path:
    file = (SKILLS / path.removeprefix("skills/")).resolve()
    if not file.is_relative_to(SKILLS.resolve()):
        raise ValueError(f"{path} is outside the skills")
    return file


def _name(workspace: Workspace, file: Path) -> str:
    return file.relative_to(workspace.root).as_posix()


def _state(workspace: Workspace) -> dict[str, Any]:
    # each file's size and time of change, by name
    return {f["name"]: (f["size"], f["modified"]) for f in workspace.files()}

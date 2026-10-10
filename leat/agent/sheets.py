"""Spreadsheets of the workspace's: tables read, of a CSV or an Excel workbook's first sheet, and
tables written as a workbook of sheets, each in the sandbox, with its libraries, as every document
is. A CSV is read whatever it is separated by, of UTF-8 or of Windows' Turkish, as offices' banks
write them.
"""

from typing import Any

from leat.agent.workspace import Workspace

# reads the table at sys.argv[2], a CSV or a workbook's first sheet, as its rows of cells, text,
# numbers or dates as ISO writes them, without the empty ones; prints them as JSON
_READ = """
import csv, datetime, json, sys
path = sys.argv[2]
if path.lower().endswith((".csv", ".tsv", ".txt")):
    data = open(path, "rb").read()
    for encoding in ("utf-8-sig", "cp1254"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            pass
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
    rows = list(csv.reader(text.splitlines(), dialect))
else:
    from openpyxl import load_workbook
    sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    rows = [[c.isoformat() if isinstance(c, (datetime.date, datetime.datetime)) else c
             for c in row] for row in sheet.iter_rows(values_only=True)]
rows = [row for row in rows if any(c not in (None, "") for c in row)]
print(json.dumps(rows, ensure_ascii=False))
"""
# writes the sheets of the JSON at sys.argv[1], each of a name and its rows, the first its
# heading, as the workbook sys.argv[2]: each heading bold, frozen and filtered, each column as wide
# as it needs
_WRITE = """
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
book = Workbook()
book.remove(book.active)
for title, rows in json.load(open(sys.argv[1], encoding="utf-8")):
    sheet = book.create_sheet(title[:31])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, width + 2))
book.save(sys.argv[2])
print("{}")
"""


def read(space: Workspace, name: str) -> list[list[Any]]:
    """A table's rows, of a CSV or a workbook's first sheet, the empty ones left out. Raises
    FileNotFoundError, or ValueError if it cannot be read."""
    if not space.path(name).is_file():
        raise FileNotFoundError(f"there is no file {name}")
    return space.given(_READ, None, name)


def write(space: Workspace, sheets: list[tuple[str, list[list[Any]]]], name: str) -> str:
    """Writes sheets, each of a title and its rows, the first its heading, as a workbook of the
    space's, by a name free there; returns the name. Raises ValueError if it cannot be written."""
    if space.environment is None:
        raise ValueError("the sandbox has no spreadsheet library")
    name = space.free(name if name.lower().endswith(".xlsx") else f"{name}.xlsx", folders=True)
    space.path(name).parent.mkdir(parents=True, exist_ok=True)
    space.given(_WRITE, sheets, name)
    return name
